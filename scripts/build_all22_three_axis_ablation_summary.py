#!/usr/bin/env python3
"""Build endpoint and aggregate tables from frozen three-axis certification."""

from __future__ import annotations

import csv
import json
import math
import statistics
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from run_classic_ml_external_data_confirmation import TASKS


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/all22_three_axis_ablation_v1"
CSV_PATH = REPORT / "all22_three_axis_ablation_results.csv"
XLSX_PATH = REPORT / "all22_three_axis_ablation_results.xlsx"
AUDIT_PATH = REPORT / "all22_three_axis_ablation_results.audit.json"
ARMS = ("baseline", "M1", "M2", "M3", "F1", "F2", "F3", "D1", "D2", "D3")
CELLS = ("000", "100", "010", "001", "110", "101", "011", "111")


def _mean_sd(values: list[float]) -> dict[str, float]:
    return {
        "mean": float(statistics.mean(values)),
        "sample_sd": float(statistics.stdev(values)) if len(values) > 1 else 0.0,
    }


def _label(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def _load(task: str) -> dict[str, Any]:
    path = REPORT / "certification" / f"{task}.json"
    error = path.with_suffix(".error.json")
    if error.exists():
        raise RuntimeError(f"Certification failed for {task}: {error}")
    if not path.exists():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not (
        payload.get("phase") == "certification"
        and payload.get("task") == task
        and payload.get("status") == "completed"
        and payload.get("safety", {}).get("test_results_used_for_selection") is False
        and payload.get("safety", {}).get("test_labels_loaded_only_after_frozen_confirmation") is True
    ):
        raise RuntimeError(f"Invalid certification artifact for {task}")
    return payload


def _effect(cells: dict[str, float], axis: int) -> float:
    deltas = []
    for cell in CELLS:
        if cell[axis] != "0":
            continue
        active = cell[:axis] + "1" + cell[axis + 1 :]
        deltas.append(cells[active] - cells[cell])
    return float(statistics.mean(deltas))


def _two_way(cells: dict[str, float], left: int, right: int) -> float:
    other = ({0, 1, 2} - {left, right}).pop()
    values = []
    for other_value in ("0", "1"):
        base = ["0", "0", "0"]
        base[other] = other_value
        c00 = "".join(base)
        base[left] = "1"
        c10 = "".join(base)
        base[left] = "0"
        base[right] = "1"
        c01 = "".join(base)
        base[left] = "1"
        c11 = "".join(base)
        values.append((cells[c11] - cells[c10]) - (cells[c01] - cells[c00]))
    return float(statistics.mean(values))


def main() -> None:
    rows: list[dict[str, Any]] = []
    task_effects: list[dict[str, float]] = []
    protocol_sha = None
    for task in TASKS:
        payload = _load(task)
        protocol_sha = protocol_sha or payload["protocol_sha256"]
        if payload["protocol_sha256"] != protocol_sha:
            raise RuntimeError("Certification artifacts use different protocols")
        ladder = payload["ladder"]["summaries"]
        factorial = payload["factorial"]["summaries"]
        route = payload["routing"]
        external = payload["external_audit"]
        selected = json.loads(
            (REPORT / "discovery" / f"{task}.json").read_text(encoding="utf-8")
        )["selected"]
        row: dict[str, Any] = {
            "Task": task,
            "Metric": payload["metric"],
            "Data Source Admitted": external["admitted"],
            "Same Source Rejected": external["source_rejected_same_source"],
            "External Rows After Firewall": external["post_firewall_rows"],
            "Routed Cell": route["selected_cell"],
            "Routed Test Score Mean": route["summary"]["mean_score"],
            "Routed Test Score SD": route["summary"]["sample_score_sd"],
            "Routed Normalized Delta": route["summary"]["mean_normalized_improvement"],
            "Selected M3": _label(selected["model"]["M3"]["config"]),
            "Selected F3": _label(selected["feature"]["F3"]["config"]),
            "Selected D3": _label(selected["data"]["D3"]["config"]),
        }
        for arm in ARMS:
            summary = ladder[arm]
            row[f"{arm} Score Mean"] = summary["mean_score"]
            row[f"{arm} Score SD"] = summary["sample_score_sd"]
            row[f"{arm} Normalized Delta"] = summary["mean_normalized_improvement"]
        cell_values = {
            cell: factorial[f"factorial_{cell}"]["mean_normalized_improvement"]
            for cell in CELLS
        }
        for cell in CELLS:
            summary = factorial[f"factorial_{cell}"]
            row[f"Cell {cell} Score Mean"] = summary["mean_score"]
            row[f"Cell {cell} Score SD"] = summary["sample_score_sd"]
            row[f"Cell {cell} Normalized Delta"] = summary["mean_normalized_improvement"]
        effects = {
            "Model Main Effect": _effect(cell_values, 0),
            "Feature Main Effect": _effect(cell_values, 1),
            "Data Main Effect": _effect(cell_values, 2),
            "Model×Feature Interaction": _two_way(cell_values, 0, 1),
            "Model×Data Interaction": _two_way(cell_values, 0, 2),
            "Feature×Data Interaction": _two_way(cell_values, 1, 2),
            "Three-way Interaction": (
                cell_values["111"]
                - cell_values["110"]
                - cell_values["101"]
                - cell_values["011"]
                + cell_values["100"]
                + cell_values["010"]
                + cell_values["001"]
                - cell_values["000"]
            ),
        }
        row.update(effects)
        task_effects.append(effects)
        rows.append(row)

    headers = list(rows[0])
    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    aggregate_effects = {
        name: _mean_sd([row[name] for row in task_effects]) for name in task_effects[0]
    }
    aggregate_arms = {
        arm: _mean_sd([row[f"{arm} Normalized Delta"] for row in rows]) for arm in ARMS
    }
    routed = _mean_sd([row["Routed Normalized Delta"] for row in rows])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Endpoint Results"
    sheet.append(headers)
    for row in rows:
        sheet.append([row[name] for name in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(displayName="ThreeAxisEndpointResults", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)

    aggregate = workbook.create_sheet("Aggregate")
    aggregate.append(["Category", "Item", "Mean", "Sample SD"])
    for arm, values in aggregate_arms.items():
        aggregate.append(["Axis ladder", arm, values["mean"], values["sample_sd"]])
    aggregate.append(["Routing", "Routed Agent", routed["mean"], routed["sample_sd"]])
    for name, values in aggregate_effects.items():
        aggregate.append(["Factorial effect", name, values["mean"], values["sample_sd"]])
    for cell in aggregate[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    aggregate.freeze_panes = "A2"
    workbook.save(XLSX_PATH)

    audit = {
        "protocol": "all22-three-axis-ablation-summary-v1",
        "source_protocol_sha256": protocol_sha,
        "tasks": len(rows),
        "all_certifications_complete": True,
        "aggregate_axis_ladders": aggregate_arms,
        "aggregate_factorial_effects": aggregate_effects,
        "routed": routed,
        "csv": str(CSV_PATH.resolve()),
        "xlsx": str(XLSX_PATH.resolve()),
    }
    AUDIT_PATH.write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n")
    print(json.dumps(audit, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
