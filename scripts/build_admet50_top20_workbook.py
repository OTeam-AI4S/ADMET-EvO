"""Build the audited 50-task screen and Top-20 confirmation Excel workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY, BLUE, TEAL, GREEN, GOLD, RED = "17365D", "2F75B5", "2F8F9D", "70AD47", "FFC000", "C00000"
WHITE, LIGHT_BLUE, LIGHT_GREEN, LIGHT_GOLD, LIGHT_RED = "FFFFFF", "D9EAF7", "E2F0D9", "FFF2CC", "FCE4D6"
GRID = Side(style="thin", color="D9E1F2")
CLASSICS = {
    "ExtraTrees-Morgan": "baseline", "ExtraTrees-rich": "feature_rich",
    "LightGBM-fixed": "boosting_anchor", "LightGBM-tuned": "automl_expanded",
    "CatBoost-GPU": "gpu_catboost", "RandomForest-rich": "random_forest",
    "LinearSVM-Morgan": "linear_svm", "XGBoost-rich": "xgboost",
}


def _native(value: float, metric: str) -> float:
    return -float(value) if metric == "mae" else float(value)


def _append_frame(ws, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))


def _style(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    if ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=GRID)


def _widths(ws) -> None:
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 18
    ws.column_dimensions["A"].width = 48


def _tables(root: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    selection_dir = root / "reports/admet50_top20_selection_strict_v2"
    summary_dir = root / "reports/admet50_top20_holdout_summary_v2"
    screen = pd.read_csv(selection_dir / "admet50_seed17_development_screen.csv")
    summary = pd.read_csv(summary_dir / "endpoint_confirmation.csv")
    folds = pd.read_csv(summary_dir / "fold_results.csv")
    cohort = pd.read_csv(root / "reports/admet50_discovery_cohort_v1.csv")
    if len(screen) != 50 or int(screen["selected_top20"].sum()) != 20 or len(summary) != 20:
        raise RuntimeError("Workbook requires an audited 50-task denominator and 20 confirmations")

    counts = cohort.groupby("endpoint_id").agg(
        identities=("inchikey", "nunique"), chemistry_groups=("chemistry_group", "nunique")
    )
    minority = cohort.loc[cohort["task_type"].eq("classification")].groupby("endpoint_id")["target"].agg(
        lambda values: int(values.value_counts().min())
    ).rename("minority_count")
    denominator = screen.merge(counts, left_on="endpoint_id", right_index=True, how="left")
    denominator = denominator.merge(minority, left_on="endpoint_id", right_index=True, how="left")
    denominator = denominator.rename(columns={
        "endpoint_id": "Endpoint", "task_type": "Task type", "metric": "Metric",
        "selected_action": "Seed17 selected action",
        "development_agent_score": "Development oriented score",
        "development_baseline": "Development baseline",
        "development_oriented_gain": "Development gain",
        "quality_percentile": "Quality percentile", "gain_percentile": "Gain percentile",
        "selection_index": "Selection index", "development_rank": "Rank / 50",
        "selected_top20": "Selected Top20", "identities": "Identities",
        "chemistry_groups": "Chemistry groups", "minority_count": "Minority count",
    })

    confirm = folds.loc[folds["confirmation_seed"].isin([29, 43])].copy()
    pivot = confirm.groupby(["endpoint_id", "action"])["oriented_score"].mean().unstack()
    rows, oriented_rows = [], []
    for record in summary.itertuples(index=False):
        metric = str(record.metric)
        agent = float(record.confirmation_score)
        classic = {label: float(pivot.loc[record.endpoint_id, action]) for label, action in CLASSICS.items()}
        best_label, best = max(classic.items(), key=lambda item: (item[1], item[0]))
        delta = agent - best
        all_scores = {"Agent": agent, **classic}
        rows.append({
            "Endpoint": record.endpoint_id, "Source": record.source_id,
            "Task type": record.task_type, "Metric": "MCC ↑" if metric == "mcc" else "MAE ↓",
            "Agent selected action": record.selected_action, "Agent score": _native(agent, metric),
            **{label: _native(value, metric) for label, value in classic.items()},
            "Holdout oracle best classic (upper bound)": best_label,
            "Holdout oracle score": _native(best, metric),
            "Agent vs holdout oracle": "胜" if delta > 1e-12 else "负" if delta < -1e-12 else "平",
            "Oriented delta vs oracle": delta,
            f"Agent rank / {len(CLASSICS)+1}": float(pd.Series(all_scores).rank(ascending=False, method="average")["Agent"]),
            "Confirmation seed SD": float(record.confirmation_seed_sd),
            "Gain vs Morgan baseline": float(record.oriented_gain),
        })
        oriented_rows.append({"Endpoint": record.endpoint_id, **all_scores})
    comparison = pd.DataFrame(rows).sort_values(["Task type", "Endpoint"], kind="stable")
    oriented = pd.DataFrame(oriented_rows).set_index("Endpoint")
    wtl = []
    for label in CLASSICS:
        delta = oriented["Agent"] - oriented[label]
        wins, ties, losses = int(delta.gt(1e-12).sum()), int(delta.abs().le(1e-12).sum()), int(delta.lt(-1e-12).sum())
        wtl.append({
            "Comparator": label, "Agent wins": wins, "Ties": ties, "Agent losses": losses,
            "Mean oriented gain": float(delta.mean()), "Median oriented gain": float(delta.median()),
        })
    ranks = oriented.rank(axis=1, ascending=False, method="average").mean().sort_values().reset_index()
    ranks.columns = ["Model", "Mean rank (lower is better)"]
    return denominator, comparison, pd.DataFrame(wtl), ranks


def build(root: Path, output: Path) -> None:
    denominator, comparison, wtl, ranks = _tables(root)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "ADMET50 Development Screen and Top20 Confirmation"
    workbook.properties.creator = "ADMET Agent"

    overview = workbook.create_sheet("Overview")
    overview.merge_cells("A1:H1")
    overview["A1"] = "50个新ADMET任务开发筛选 → Top-20独立确认"
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].font = Font(color=WHITE, bold=True, size=16)
    overview["A1"].alignment = Alignment(horizontal="center")
    overview.append([])
    overview.append(["Protocol item", "Value"])
    overview.append(["Full denominator", "50 endpoints; every task remains visible in this workbook"])
    overview.append(["Development selection", "Seed 17, three-fold CV within folds 0–2 only; quality/gain percentile rule"])
    overview.append(["Top-20 confirmation", "Fit on folds 0–2; one untouched holdout = union of folds 3–4; seeds 29/43 test fit stability"])
    overview.append(["Split", "46,975 holdout rows; chemistry-group and identity overlap with model-selection partition are both zero"])
    overview.append(["Agent warm starts", "ExtraTrees, LightGBM, CatBoost, RandomForest, linear SVM, XGBoost action library"])
    overview.append(["Metrics", "Classification MCC ↑; regression MAE ↓; oriented deltas are positive when Agent is better"])
    overview.append(["Top20 improved vs Morgan", f"{int((comparison['Gain vs Morgan baseline'] > 0).sum())} / 20"])
    for label in ["CatBoost-GPU", "RandomForest-rich", "XGBoost-rich"]:
        row = wtl.loc[wtl["Comparator"].eq(label)].iloc[0]
        overview.append([
            f"Agent vs {label}",
            f"{int(row['Agent wins'])} wins / {int(row['Ties'])} ties / {int(row['Agent losses'])} losses",
        ])
    for cell in overview[3]:
        cell.fill = PatternFill("solid", fgColor=TEAL); cell.font = Font(color=WHITE, bold=True)
    for row in overview.iter_rows(min_row=4, max_row=overview.max_row):
        row[0].font = Font(bold=True, color=NAVY); row[1].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    overview.column_dimensions["A"].width = 30; overview.column_dimensions["B"].width = 100

    full = workbook.create_sheet("Full 50 Development")
    _append_frame(full, denominator); _style(full); _widths(full)
    h = {cell.value: cell.column for cell in full[1]}
    for row in range(2, full.max_row + 1):
        if full.cell(row, h["Selected Top20"]).value:
            full.cell(row, h["Selected Top20"]).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    for name in ["Development gain", "Selection index"]:
        col = get_column_letter(h[name])
        full.conditional_formatting.add(
            f"{col}2:{col}{full.max_row}",
            ColorScaleRule(start_type="min", start_color=RED, mid_type="percentile", mid_value=50,
                           mid_color=LIGHT_GOLD, end_type="max", end_color=GREEN),
        )

    top = workbook.create_sheet("Top20 Confirmation")
    _append_frame(top, comparison); _style(top); _widths(top)
    htop = {cell.value: cell.column for cell in top[1]}
    for row in range(2, top.max_row + 1):
        cell = top.cell(row, htop["Agent vs holdout oracle"])
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN if cell.value == "胜" else LIGHT_RED if cell.value == "负" else LIGHT_GOLD)

    wtl_sheet = workbook.create_sheet("Win-Tie-Loss")
    _append_frame(wtl_sheet, wtl); _style(wtl_sheet); _widths(wtl_sheet)
    wtl_chart = BarChart(); wtl_chart.type = "bar"; wtl_chart.grouping = "stacked"; wtl_chart.overlap = 100
    wtl_chart.title = "Top20: Agent vs each classic model"; wtl_chart.x_axis.title = "Endpoint count"
    wtl_chart.add_data(Reference(wtl_sheet, min_col=2, max_col=4, min_row=1, max_row=wtl_sheet.max_row), titles_from_data=True)
    wtl_chart.set_categories(Reference(wtl_sheet, min_col=1, min_row=2, max_row=wtl_sheet.max_row))
    wtl_chart.height = 9; wtl_chart.width = 16; wtl_sheet.add_chart(wtl_chart, "H2")

    chart_data = workbook.create_sheet("Chart Data")
    dev = denominator.sort_values("Rank / 50")[["Endpoint", "Development gain", "Selection index", "Selected Top20"]]
    _append_frame(chart_data, dev)
    start_top = chart_data.max_column + 3
    top_chart = comparison.loc[comparison["Task type"].eq("classification"), [
        "Endpoint", "Agent score", "Holdout oracle score", "Metric"
    ]]
    for offset, name in enumerate(top_chart.columns): chart_data.cell(1, start_top + offset, name)
    for row_index, row in enumerate(top_chart.itertuples(index=False, name=None), start=2):
        for offset, value in enumerate(row): chart_data.cell(row_index, start_top + offset, value)
    start_reg = start_top + len(top_chart.columns) + 2
    reg_chart_data = comparison.loc[comparison["Task type"].eq("regression"), [
        "Endpoint", "Agent score", "Holdout oracle score", "Metric"
    ]]
    for offset, name in enumerate(reg_chart_data.columns): chart_data.cell(1, start_reg + offset, name)
    for row_index, row in enumerate(reg_chart_data.itertuples(index=False, name=None), start=2):
        for offset, value in enumerate(row): chart_data.cell(row_index, start_reg + offset, value)
    start_rank = start_reg + len(reg_chart_data.columns) + 2
    for offset, name in enumerate(ranks.columns): chart_data.cell(1, start_rank + offset, name)
    for row_index, row in enumerate(ranks.itertuples(index=False, name=None), start=2):
        chart_data.cell(row_index, start_rank, row[0]); chart_data.cell(row_index, start_rank + 1, float(row[1]))
    _style(chart_data); _widths(chart_data)

    charts = workbook.create_sheet("Numeric Charts")
    charts["A1"] = "ADMET50 / Top20 numerical comparison"; charts["A1"].font = Font(size=16, bold=True, color=NAVY)
    dev_chart = BarChart(); dev_chart.type = "bar"; dev_chart.title = "All 50: seed17 oriented gain over Morgan baseline"
    dev_chart.x_axis.title = "Oriented gain (positive is better)"; dev_chart.height = 25; dev_chart.width = 24
    dev_chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=51), titles_from_data=True)
    dev_chart.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=51)); charts.add_chart(dev_chart, "A3")
    top_score = BarChart(); top_score.type = "bar"; top_score.grouping = "clustered"
    top_score.title = "Top20 classification: Agent vs post-hoc classic oracle upper bound (MCC)"; top_score.height = 15; top_score.width = 23
    top_score.add_data(Reference(chart_data, min_col=start_top+1, max_col=start_top+2, min_row=1, max_row=1+len(top_chart)), titles_from_data=True)
    top_score.set_categories(Reference(chart_data, min_col=start_top, min_row=2, max_row=1+len(top_chart))); top_score.dLbls = DataLabelList(); top_score.dLbls.showVal = True
    charts.add_chart(top_score, "X3")
    if len(reg_chart_data):
        regression_chart = BarChart(); regression_chart.type = "col"; regression_chart.grouping = "clustered"
        regression_chart.title = "Top20 regression: Agent vs post-hoc classic oracle upper bound (MAE)"; regression_chart.height = 8; regression_chart.width = 16
        regression_chart.add_data(Reference(chart_data, min_col=start_reg+1, max_col=start_reg+2, min_row=1, max_row=1+len(reg_chart_data)), titles_from_data=True)
        regression_chart.set_categories(Reference(chart_data, min_col=start_reg, min_row=2, max_row=1+len(reg_chart_data)))
        regression_chart.dLbls = DataLabelList(); regression_chart.dLbls.showVal = True; charts.add_chart(regression_chart, "X34")
    rank_chart = BarChart(); rank_chart.type = "bar"; rank_chart.title = "Top20 mean model rank"
    rank_chart.add_data(Reference(chart_data, min_col=start_rank+1, min_row=1, max_row=1+len(ranks)), titles_from_data=True)
    rank_chart.set_categories(Reference(chart_data, min_col=start_rank, min_row=2, max_row=1+len(ranks))); rank_chart.dLbls = DataLabelList(); rank_chart.dLbls.showVal = True
    rank_chart.height = 9; rank_chart.width = 16; charts.add_chart(rank_chart, "X51")

    notes = pd.DataFrame([
        ["Endpoint pool", "17 frozen public endpoints + 20 ToxCast + 13 SIDER; source admission did not inspect model outcomes."],
        ["Top20", "Selected by seed17 CV confined to folds0–2; the full 50-task denominator remains visible."],
        ["Action selection", "Best mean oriented three-fold action inside folds0–2; folds3–4 are absent from every selection fit and score."],
        ["Warm start", "Reuse means a fixed model/action library, not pretrained weights or access to confirmation labels."],
        ["Fair comparison", "All actions train on folds0–2 and predict the same untouched fold3–4 holdout for each endpoint."],
        ["Oracle column", "The per-endpoint best classic is selected after viewing holdout scores and is therefore an unattainable upper bound, not a fair deployed comparator; use Win-Tie-Loss for matched individual-model comparisons."],
        ["Interpretation", "Matched comparisons support claims for this protocol; unrelated literature splits remain contextual only."],
    ], columns=["Item", "Definition / constraint"])
    note_sheet = workbook.create_sheet("Methods & Notes"); _append_frame(note_sheet, notes); _style(note_sheet); _widths(note_sheet)
    note_sheet.column_dimensions["B"].width = 110

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path("."))
    parser.add_argument("--output", type=Path, default=Path("reports/admet50_top20_model_comparison_2026-08-27.xlsx"))
    args = parser.parse_args(); build(args.root.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
