#!/usr/bin/env python3
"""Build the ADMET-EvO Results evidence workbook and master CSV.

The workbook separates publishable evidence, qualified evidence, negative
results, and incomplete work. It also imports the existing task-level tables
without changing their source files.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path("/home/wyilin/admet")
REPORTS = ROOT / "reports"
OUT_XLSX = REPORTS / "ADMET_EvO_Results数据总表_2026-08-25.xlsx"
OUT_CSV = REPORTS / "ADMET_EvO_Results证据主表_2026-08-25.csv"

MASTER_HEADERS = [
    "Result",
    "证据ID",
    "建议小标题",
    "核心问题/主张",
    "数据集/对象",
    "实验设计/样本量",
    "主要指标",
    "当前结果",
    "统计阶段",
    "证据状态",
    "建议写法",
    "不可写/限制",
    "来源文件",
    "优先级",
]


def source(path: str) -> str:
    return str(ROOT / path)


MASTER_ROWS = [
    ["R1", "R1-E1", "Endpoint contracts define the research object", "22 个任务是否在研究前被完整定义", "TDC ADMET-22", "22 endpoints × 13 核心合同字段", "合同覆盖率/字段完整率", "22/22 合同；13 个核心字段完整率 100%", "discovery audit", "主文可用", "写成任务边界与数据准入的量化审计结果", "不要把合同完整率写成预测性能提升", source("reports/admet_evo_five_results_experiment_by_experiment_v3_2026-08-24.md"), "P1"],
    ["R1", "R1-E2", "Hardness profiles are endpoint-specific", "HardnessCard 是否产生异质诊断", "TDC ADMET-22", "22/22 discovery-only diagnosis packets；test_loaded=false", "hardness/路由分布", "mean hardness 0.313；17 feature、4 model、1 data；mean domain-shift severity 0.748", "discovery audit", "主文可用（需限定）", "强调 endpoint-specific profile 与动作候选差异", "尚无 multi-seed ICC/kappa；不要称 hardness 已稳定或有外部效度", source("reports/v3_tdc22_hardness_summary.json"), "P0"],
    ["R1", "R1-E3", "Hypotheses can be supported, refuted or deferred", "可证伪 action lineage 是否保留正/负/不确定证据", "BBB、CYP2C9 Veith、CYP2D6 Veith", "matched interventions；5 seeds（CYP2C9 完成 4 对）", "oriented delta/seed 方向", "BBB +0.007006（4/5 正，supported）；CYP2C9 −0.000889+timeout（inconclusive）；CYP2D6 −0.009210（0/5 正，refuted）", "validation-only", "主文可用（需限定）", "作为闭环能处理三种证据状态的代表性谱系", "三个例子不能代表 22-task policy superiority", source("reports/admet_evo_five_results_experiment_by_experiment_v3_2026-08-24.md"), "P1"],
    ["R1", "R1-E4", "The evidence ledger retains all outcomes", "研究过程是否可追踪并保留失败", "production research ledger", "834 hash-chained records", "记录数/链验证", "789 completed、1 failed、22 plans、22 action evaluations；hash chain valid", "process audit", "主文可用", "量化写出全结果留存和失败留存", "记录数证明可审计，不证明决策科学正确", source("reports/research_process_baseline_v3.json"), "P1"],
    ["R1", "R1-E5", "Protocol and tamper controls fail closed", "关键泄漏、篡改和 source-firewall 是否被阻止", "regression test suite", "冻结证据截点的完整套件", "测试通过数", "570 passed、1 skipped；覆盖 ledger/protocol/source/test/freeze invariants", "software/process audit", "主文可用（需限定）", "写为冻结版本的回归测试结果", "当前工作树有后续改动；不要将该数字暗示为今天代码的重新验证结果", source("reports/admet_evo_results_paper_framed_2026-08-23.md"), "P1"],
    ["R1", "R1-E6a", "Finalists are frozen before certification", "TDC finalist 是否内容寻址并完整覆盖", "TDC ADMET-22", "22/22 frozen finalists；至少 5 finalization seeds", "freeze coverage/SHA-256", "22/22；bundle SHA-256 693db693…8f9e", "application freeze", "主文可用（需限定）", "写为 application qualification 的冻结边界", "TDC 历史 test 已接触，不能称全新 prospective blind certification", source("reports/goal_scorecard_v27.json"), "P0"],
    ["R1", "R1-E6b", "Prospective artifacts are content-addressed", "OpenADMET recipe 是否在 blind feedback 前冻结", "OpenADMET CYP", "B1 recipe + 750-row blind inputs/predictions", "freeze/feedback flags", "recipe SHA-256 2182b275…d1e；blind labels loaded=false；selection reopened=false", "prospective blind-inference freeze", "主文可用（需限定）", "强调决策与 artifact 在外部结果前冻结", "不能写已经获得官方 blind 性能或排名", source("reports/openadmet_cyp_b1_prospective_freeze_v1.md"), "P0"],

    ["R2", "R2-E1a", "Broad performance across heterogeneous endpoints", "与来源可追溯的模型相比是否具备跨任务竞争力", "TDC ADMET-22；8-model author/source table", "22 tasks；8 systems/models；five-run means", "within-table mean rank/best count", "ADMET-EvO mean rank 2.977（第 2/8）；best 4/22；MapLight+GNN 2.500（第 1/8）", "frozen benchmark application scores", "主文可用", "把 8-model 来源对照作为主性能表", "不是 live leaderboard；不同模型年代/预训练暴露不同", source("reports/model_comparison_tdc22_8models_2026-08-25.csv"), "P0"],
    ["R2", "R2-E1b", "Captured-board sensitivity analysis", "冻结分数插入历史页面后的估计覆盖", "captured TDC pages 2026-08-17", "22 endpoints；snapshot insertion", "estimated rank counts", "Top-1 5/22、Top-5 18/22、Top-10 21/22；mean rank 4.545；median 4.0", "historical snapshot estimate", "补充材料/需限定", "作为 sensitivity/application qualification，明确 estimated", "不能称官方 Top-1/Top-5；Bioavailability 页面存在异常", source("reports/goal_scorecard_v27.json"), "P0"],
    ["R2", "R2-E2", "Comparison with complete-coverage fixed models", "是否超过多数完整覆盖固定模型", "captured leaderboard；14 fixed models", "只纳入 22/22 同名完整覆盖模型", "combined rank/W-T-L", "ADMET-EvO 第 2/15；胜 13/14；vs MapLight 13/0/9，vs Chemprop-RDKit 20/0/2，vs AttentiveFP 21/0/1", "historical snapshot analysis", "补充材料/需限定", "保留为完整覆盖 sensitivity analysis", "MapLight+GNN 仍更优；受网页快照完整性影响", source("reports/cross_endpoint_fixed_model_comparison_v1.json"), "P1"],
    ["R2", "R2-E3a", "Targeted graph expertise can help selected tasks", "GNN expert 在 Bioavailability 内部配对验证是否有增益", "Bioavailability_Ma", "5 paired split seeds 6–10；validation only", "AUROC oriented delta", "+0.023526；4/5 wins；paired-t 95% CI [0.000137, 0.046915]；sign-flip p=0.125", "post-hoc paired validation", "主文可用（需限定）", "用 suggested/improved in paired validation；报告 n 与区间", "不能与 author-source benchmark test score混写；不是独立 endpoint replicate", source("reports/primary_model_effects_v27.json"), "P1"],
    ["R2", "R2-E3b", "Targeted graph expertise can help selected tasks", "GNN expert 在 CYP2C9 substrate 内部配对验证是否有增益", "CYP2C9_Substrate_CarbonMangels", "5 paired split seeds 6–10；validation only", "AUPRC oriented delta", "+0.034225；4/5 wins；95% CI [−0.042403, 0.110852]；sign-flip p=0.1875", "post-hoc paired validation", "主文可用（需弱化）", "写为方向性/探索性增益", "CI 跨 0；不能写 reproduced significant improvement", source("reports/primary_model_effects_v27.json"), "P1"],
    ["R2", "R2-E3c", "Model scaling is not uniformly beneficial", "扩大 UniMol2 是否稳定改善 Lipophilicity", "Lipophilicity_AstraZeneca", "5 paired seeds；164M vs 84M", "MAE-oriented delta/runtime ratio", "mean −0.009342；1/5 wins；95% CI [−0.020574, 0.001889]；runtime 2.03×", "post-hoc paired validation", "主文可用（负结果）", "作为选择性 acquisition 而非统一扩模的反例", "不能称 84M 普遍优于 164M", source("reports/primary_model_effects_v27.json"), "P1"],
    ["R2", "R2-E4", "Policy value is not yet established against AutoML", "当前 Agent policy 是否对 expanded AutoML 非劣", "TDC ADMET-22", "22 tasks；三臂 fresh validation；154 logical fits", "normalized delta/bootstrap/W-T-L", "Agent−expanded AutoML −0.0098；95% CI [−0.0277, 0.0065]；W/T/L 5/7/10；noninferiority FAIL", "fresh validation control", "负结果（必须披露）", "放在 R2 末尾或 Limitations，明确 broad model performance 不等于 policy superiority", "不能写 hardness-aware policy improves prediction or beats AutoML", source("reports/minimal_predictive_control_v1.json"), "P0"],
    ["R2", "R2-E5", "Compute-aware selection is practically noninferior", "cost-aware 选择能否保持验证性能并减少选中 recipe 成本", "TDC ADMET-22", "22 fresh-validation tasks；seeds 4–5；margin −0.005", "mean delta/bootstrap/cost ratio", "mean −0.000348；95% CI [−0.001482, 0.000585]；20/22 nonmaterial；cost ratio 0.278；estimated reduction 72.2%", "fresh validation efficiency", "主文可用（需限定）", "写为 selected-recipe fitting cost 的 practical noninferiority", "不是端到端 wall-clock/GPU/LLM 成本；节省主要由 LD50 昂贵 recipe 驱动", source("reports/practical_efficiency_certification_v1.json"), "P0"],
    ["R2", "R2-E6a", "Full-outcome memory mainly reduces cost", "完整 outcome memory 是否提高平均性能", "TDC ADMET-22", "leave-one-task-out validation controls", "mean gain/bootstrap", "full-outcome−winner-only −0.00031；95% CI [−0.00276, 0.00240]；superiority FAIL；compute proxy 18.5%", "validation-only", "负结果（必须披露）", "把价值限定为成本/审计，不写准确率优势", "不能称 evidence memory improves prediction", source("reports/compute_memory_appendix_v27.json"), "P1"],
    ["R2", "R2-E6b", "Matched-cost random control", "action selection 是否优于同成本随机", "TDC ADMET-22", "matched-cost random control", "mean gain/bootstrap", "Agent−random +0.00083；95% CI [−0.00380, 0.00542]；superiority FAIL", "validation-only", "负结果（必须披露）", "作为当前 router 局限", "不能称 action selection 已建立优越性", source("reports/compute_memory_appendix_v27.json"), "P1"],
    ["R2", "R2-E7", "Error increases outside familiar chemistry", "低相似度化学空间是否风险更高", "13 classification + 9 regression endpoints", "validation calibration/residual audit", "Brier/ECE/low-similarity error", "mean Brier 0.1173；mean adaptive ECE 0.0964；11/13 分类 Brier 更差；7/9 回归 residual 更高", "validation-side applicability audit", "主文可用（需限定）", "支持 applicability diagnosis 的必要性", "不是 prospective test OOD certification", source("reports/validation_calibration_residuals_v27.json"), "P1"],

    ["R3", "R3-E1a", "Action effects are endpoint-dependent", "Feature action 在完整反事实网格的效应", "5 sentinel endpoints", "5 endpoints × 3 actions × 5 seeds；100 fits", "supported/refuted/posterior", "Feature 3 supported、2 refuted；posterior support 0.556", "prospective validation action grid", "主文可用", "强调异质性，不只摘录正增益", "不能称 Feature 对所有 endpoint 均有效", source("reports/transport_grid_results_v1.json"), "P0"],
    ["R3", "R3-E1b", "Action effects are endpoint-dependent", "Model action 在完整反事实网格的效应", "5 sentinel endpoints", "RandomForest vs ExtraTrees；5 seeds/endpoint", "supported/refuted/posterior", "4 refuted、1 inconclusive；posterior support 0.278", "prospective validation action grid", "主文可用（负结果）", "展示 model-family prior 不可全局推广", "不能只写图模型正例来概括 Model axis", source("reports/transport_grid_results_v1.json"), "P0"],
    ["R3", "R3-E1c", "Action effects are endpoint-dependent", "Generic data filtering 是否产生可用先验", "5 sentinel endpoints", "quality filter；5 matched episodes", "evidence class", "5/5 inconclusive；改变 25/25 splits 的训练成员；不进入 routing prior", "prospective validation action grid", "主文可用（负结果）", "表明数据被改变不等于产生可靠增益", "不能写 generic data action improved prediction", source("reports/transport_grid_results_v1.json"), "P1"],
    ["R3", "R3-E2", "Independent confirmation can overturn memory", "负向 model prior 是否可安全剪枝", "5 fresh endpoints", "50/50 fits；fresh seeds", "supported/refuted/pruning regret", "2 supported、2 refuted、1 inconclusive；mean positive pruning regret 0.065379；max 0.258620；global pruning rolled back", "independent post-memory confirmation", "主文可用", "这是最直接的 self-correction/self-evolution 证据之一", "不要写 memory prior 已提升总体准确率", source("reports/post_memory_confirmation_v1.json"), "P0"],
    ["R3", "R3-E3", "Early EC-VOI routing exposed calibration failure", "v1 router 与 legacy 比较", "3 development endpoints", "EC-VOI vs legacy/rules/random", "W-T-L/delta/cost", "79/80 fits；1 timeout；vs legacy 0/1/2，mean −0.009598；cost 3.18×", "development", "负结果（必须披露）", "作为算法迭代的 failure localization", "不能作为 agent superiority 证据", source("reports/ecvoi_implementation_status_v1.json"), "P1"],
    ["R3", "R3-E4", "Shared anchors remove weak-baseline confounding", "v2/v3 是否超越共享强锚点", "3 development endpoints", "v2 13 arms；v3 7 arms", "promotion gates", "v2 vs AutoML 0/0/3，mean −0.032686；v3 nonloss 但 0/3 严格超越 anchor；两代均 FAIL", "development", "负结果（必须披露）", "说明 anchor 修复混杂但未建立 router 增益", "19-task confirmation 未解封", source("reports/ecvoi_v3_development_gate_v3.json"), "P1"],
    ["R3", "R3-E5", "Diagnosis matters only under constrained allocation", "matched-state diagnosis 的可识别价值", "3 endpoints × 5 states", "3 actions × 3 seeds = 135 fits", "gain/regret/harm", "unconstrained 时所有策略 15/15 选 rich-2D；50% budget 下 Full gain 0.068074，oracle regret 低 34.1%；1 harmful intervention；gate FAIL", "matched-state development", "主文可用（需限定）", "强调预算条件与机制 gate 失败", "不能写 diagnosis uniformly outperforms no-diagnosis", source("reports/ecvoi_v4_matched_state_full_analysis_v1.json"), "P0"],
    ["R3", "R3-E6", "Cross-source routing did not outperform simple policies", "v5 source-out action selection 是否通过", "4 sources；16 endpoints；415 rows", "Full/no-diagnosis/rules/fixed/AutoML", "gain/regret/harm/cost", "Full gain 0.076882、regret 0.045966、harm 0、cost 45；fixed LGBM gain 0.125995；promotion FAIL", "public cross-source screen", "负结果（必须披露）", "把 zero-harm 与 selection failure 同时写出", "不能称 diagnosis-conditioned transfer 已验证", source("reports/ecvoi_v5_public_policy_analysis_v1.json"), "P0"],
    ["R3", "R3-E7", "Identifiability guard abstains when diagnosis is unsupported", "v6 是否形成新来源确认", "v5 replay + frozen new sources", "16 files ≈198 MB；未运行模型", "identifiability/action flags", "identifiable states 0；max shrinkage 0；回退 no-diagnosis；model_actions_run=false", "development replay/data freeze", "未完成/不可宣称", "可以写 guard 拒绝伪外推", "不能写 fresh-source confirmation 已完成", source("reports/ecvoi_v6_v5_development_replay_v1.json"), "P1"],
    ["R3", "R3-E8a", "Adaptive external-data routing across 22 tasks", "最新 v1 数据动作探索筛出了哪些任务", "TDC ADMET-22 + curated external cohorts", "development-only；任务特异 LightGBM action", "augment/abstain/generalization count", "9/22 augment、13/22 abstain；9 个开发增益均 >0.005；8/9 在至少一个 classic model 上同向", "development-only", "探索性/待独立确认", "作为新增数据动作 screening，放 Supplementary/未来工作", "混合指标不可直接平均；不是 frozen optimal-recipe confirmation", source("reports/all22_adaptive_agent_data_v1/all22_data_augmentation_summary.csv"), "P0"],
    ["R3", "R3-E8b", "Independent optimal-recipe data confirmation", "Caco2 数据动作能否在冻结最优 recipe 与新 seeds 上复现", "Caco2_Wang + ChEMBL37 cohort", "frozen LightGBM；seeds 8–11；same recipe both arms", "MAE-oriented delta/promotion", "+0.010208 ±0.018126；3/4 seeds positive；promotion passed；test_loaded=false", "independent validation confirmation", "待来源许可/lineage 清理", "清理 source lineage 后可升级为 R3 主文数据动作案例", "当前 lineage_status=pending_document_lineage_clearance；不能立即当可发表外部数据证据", source("reports/all22_optimal_recipe_adaptive_data_v2/confirmation/Caco2_Wang.json"), "P0"],
    ["R3", "R3-GAP", "System-level axis ablation", "五轴是否各自贡献必要能力", "预注册 endpoint cohort", "Full 与 −Feature/−Model/−Data/−Replicate/−Stop；matched max budget + realized cost", "effect/CI/harm/evidence yield", "尚未执行", "planned", "未完成/不可宣称", "若保留 Result 3 的五轴总论，这是最高优先级缺口", "当前只能用分散案例说明各轴角色，不能宣称系统级互补性已证明", source("reports/admet_evo_five_results_experiment_by_experiment_v3_2026-08-24.md"), "P0"],

    ["R4", "R4-E1", "Prespecified hidden faults are detected", "结构化 guard 能否识别常见 endpoint/data 故障", "10-case commit/reveal fixture", "1 clean + 9 faults", "precision/recall/F1/critical recall/remediation", "Agent 全为 1.0；no-hardness F1=0；expanded AutoML F1=0.615；complete rules 同为 1.0", "local deterministic sealed fixture", "主文可用（需限定）", "标题用 prespecified failure modes；明确与完整规则持平", "不是独立专家 unseen benchmark；不能称超过 expert rules", source("reports/v3_hidden_fault_suite.json"), "P0"],
    ["R4", "R4-E2", "Data stewardship prevents unsafe merges", "数据准入是否降低 unsafe merge", "11 sealed data dossiers", "Agent/rules/schema/direct-fit/no-data/random", "accuracy/F1/unsafe merges", "Agent accuracy/F1/manifest=1.0，unsafe=0；expanded AutoML unsafe=7；direct-fit unsafe=10；complete rules 持平", "local sealed control", "主文可用（需限定）", "写结构化 contract steward 的安全收益和 reject-all 机会损失", "不是外部独立盲评", source("reports/data_stewardship_controls_v1.json"), "P0"],
    ["R4", "R4-E3", "Failed hypotheses trigger bounded repair and stop", "被否定/不完整证据后是否执行有限修复", "CYP2C9 substrate + hERG", "2/2 scheduled；12 fits；67.6 s", "child delta/stop", "−0.016534 与 −0.003124，均 inconclusive；两条路径均停止；无 HPO/ensemble", "prospective validation-only smoke", "主文可用（需限定）", "作为安全停止与跨轴修复案例", "2 个 case 不能推断普遍有效", source("reports/self_correction_smoke_v1.json"), "P1"],
    ["R4", "R4-E4", "The Agent rolls back a harmful learned policy", "是否能否定历史 memory 的有害剪枝", "5 fresh-validation endpoints", "独立 post-memory confirmation", "caught regret/rollback", "2 个稳定正增益本会被错剪；max caught regret 0.258620；negative_gate_confirmed=false；action restored", "independent validation confirmation", "主文可用", "与 R3-E2 合并作为最强 policy self-correction 结果", "不能写 rollback 提升了 blind test performance", source("reports/post_memory_confirmation_v1.json"), "P0"],
    ["R4", "R4-E5", "Composite research decisions integrate safety and cost", "统一 policy 能否处理复合情境", "10 sealed composite episodes", "最多 1 follow-up/6 estimated fits", "decision/process/unsafe/fits", "Agent macro=1.0、unsafe=0、fits=30；rules macro=0.86、unsafe=3、fits=48；delta +0.14 [0.04,0.26]", "local sealed challenge", "主文可用（需限定）", "写 local sealed process challenge", "不是 independent external blind validation", source("reports/composite_research_challenge_v1.json"), "P1"],
    ["R4", "R4-E6", "Committee activation is conditional", "多角色审核是否避免默认委员会开销", "22 historical states", "single/targeted/full replay", "activation/calls/fits", "20/22 single；2/22 escalate；避免 90.9% committee calls、92.4% role assignments；新增 fits=0", "behavior replay", "补充材料", "作为系统效率/异常处理附录", "未证明预测增益；不应占据 R4 主叙事", source("reports/conditional_committee_activation_v1.json"), "P2"],
    ["R4", "R4-E7", "Applicability signals identify high-risk chemistry", "UQ/距离能否识别高误差样本", "OpenADMET 6 endpoints", "10,219 records；5 folds；3 seeds", "high-error AUROC/coverage", "combined uncertainty 5/6 AUROC>0.5；low-similarity quartile 5/6 更差；回归 90% coverage≈0.898–0.900；combined gate 0/6", "internal cross-validation", "主文可用（需限定）", "移到 R5 applicability 或 Extended Data 更顺畅", "combined UQ promotion 失败；不能称统一 UQ 已验证", source("reports/openadmet_cyp_applicability_analysis_v3.json"), "P1"],
    ["R4", "R4-GAP", "Independent expert blind fault pack", "安全结果能否跨出自编 fixture", "unseen composite tasks", "至少 12 cases；4 clean + 8 faulty；双人盲评", "detection/root cause/repair/abstain", "尚无返回结果", "external blind", "未完成/不可宣称", "若 Result 4 保留强标题，这是决定性缺口", "没有独立 pack 前标题不要写 broadly detects scientific failure modes", source("reports/expert_blind_scientific_task_protocol_v1.md"), "P0"],

    ["R5", "R5-E1", "A multi-fidelity CYP study starts from endpoint contracts", "CYP 数据是否按 assay/label fidelity 拆分", "OpenADMET CYP", "Direct DRC 4,905；TDI 6,145；Emax 6,146；single concentration 17,504；blind 750", "rows/valid labels/prevalence", "4 regression labels 1,412/1,285/1,493/2,335；TDI positive 21.6%/21.3%；blind 无 label", "prospective data intake", "主文可用", "强调 known endpoints、unseen blind molecules、多 fidelity 合同", "不要包装成 unknown endpoint；endpoint/schema 是已知的", source("reports/openadmet_cyp_data_manifest_v1.json"), "P0"],
    ["R5", "R5-E2", "Identity-locked chemistry groups control leakage", "分割是否阻止分子/化学系列泄漏", "OpenADMET CYP", "6,145 identities；3,869 groups；5 outer folds；sealed 12.51%", "overlap gates", "identity/group overlap PASS；blind labels unread", "split audit", "主文可用", "作为真实 case study 的验证设计核心", "内部 certification 仍不是外部 blind readout", source("reports/openadmet_cyp_identity_locked_splits_v1.manifest.json"), "P0"],
    ["R5", "R5-E3", "The workflow rejects unnecessary complexity", "Stage B 是否选择复杂模型", "OpenADMET CYP", "B1/M3/M4；5 folds × 3 seeds", "ST-RAE/MCC", "B1 0.72594/0.24635；M3 0.77439/0.11544；M4 0.90255/0.10740；M3/M4 rejected", "development model selection", "主文可用", "写为 blind 前否证复杂方案并冻结单组件 B1", "不能称复杂表示在所有 CYP 无效", source("reports/openadmet_cyp_stage_b_component_gate_v1.json"), "P0"],
    ["R5", "R5-E4", "External evidence is routed by task type", "29,008 外部 CYP records 是否产生正/负迁移", "TDC/Veith bridge → OpenADMET CYP", "identity firewall；4 regression + 2 TDI classification", "fold-matched treatment gain", "4 regression ST-RAE gain +0.001400 至 +0.007864；2 TDI MCC gain −0.007959/−0.003875；regression retain，classification revert", "development CV", "主文可用（需限定）", "这是 Data Action 条件性最强案例；写全 6 个方向", "不能称外部数据普遍有效或已提升 blind score", source("reports/openadmet_external_propensity_gate_v2.json"), "P0"],
    ["R5", "R5-E5", "Frozen B1 passes one-time internal certification", "冻结 recipe 是否通过预注册内部 gate", "OpenADMET CYP", "sealed 12.51% internal certification", "macro/worst ST-RAE & MCC", "macro ST-RAE 0.684642≤0.825942；macro MCC 0.276628≥0.146346；worst ST-RAE 0.947342≤1.25；worst MCC 0.115551≥0；4/4 PASS", "one-time internal certification", "主文可用（需限定）", "明确是内部一次性 certification", "不能称 external blind performance", source("reports/openadmet_cyp_b1_certification_protocol_v1.json"), "P0"],
    ["R5", "R5-E6", "Blind predictions are frozen before feedback", "是否生成完整且不可变的 blind prediction", "750 blind molecules", "4 regression columns + 2 classification columns", "coverage/schema/hashes", "两份均 750 rows；schema PASS；reg SHA 53d57a28…37ae；cls SHA 614dc445…7231；feedback=false", "prospective blind-inference freeze", "主文可用（需限定）", "标题用 before external evaluation / frozen blind inference", "官方 score/rank 尚未完成，不能写 demonstrates blind performance", source("reports/openadmet_cyp_b1_prospective_freeze_v1.md"), "P0"],
    ["R5", "R5-E7", "External-source failures terminate safely", "hERG source 获取失败时是否接受不完整数据", "3 hERG candidate sources；2 ChEMBL attempts", "endpoint/license/provenance gate", "disposition/fits", "3/3 bridge_required；2/2 HTTP 500 failed closed；0 partial accept；0 modeling fit", "source acquisition safety case", "补充材料", "作为 fail-closed real-world vignette", "没有 predictive-gain result，不要放在性能主张中", source("reports/chembl_herg_external_ticket_v1.json"), "P2"],
    ["R5", "R5-GAP", "External blind readout", "冻结 workflow 是否在真正外部结果上有效", "OpenADMET blind evaluation", "官方提交/readout", "task scores/rank/CV→blind gap", "尚未完成/未归档", "external blind", "未完成/不可宣称", "这是 Result 5 唯一决定性缺口；结果返回后锁定一次分析", "不要用内部 certification 替代外部 blind outcome", source("reports/FINAL_EVIDENCE_INDEX.md"), "P0"],
]


REWRITE_HEADERS = ["Result", "当前标题/问题", "建议标题", "结构调整", "必须加入", "必须弱化或移出"]
REWRITE_ROWS = [
    ["R1", "过多 Methods 叙述；self-evolving 证据未在本节量化", "ADMET-EvO establishes an auditable, hardness-aware autonomous research loop", "把 Endpoint Contract、HardnessCard、ledger、freeze 的方法细节移至 Methods；正文用 1 个审计总表 + 1 个真实 lineage", "22/22 contracts、hardness 路由、834 ledger、570 tests、22/22 freeze", "在缺少 longitudinal external confirmation 前，不把 self-evolving 作为未经限定的首要结果"],
    ["R2", "历史 leaderboard 数字与来源对照、效率和 expert acquisition 混在一起", "ADMET-EvO is broadly competitive across 22 endpoints while reducing selected-recipe compute", "8-model author/source comparison 为主表；captured-board rank 放 sensitivity；效率单独一段", "第 2/8、mean rank 2.977、best 4/22；效率 −0.000348 与 72.2%；Agent-vs-AutoML FAIL", "Top-1=5 只能称 snapshot estimate；72.2% 不是端到端成本；GNN 结果必须标 validation/post-hoc"],
    ["R3", "当前只展示正例，容易被认为 cherry-picking", "Controlled interventions reveal endpoint-dependent action utility and trigger policy rollback", "按完整 action grid→独立复核→跨源失败→最新 adaptive-data confirmation 排序", "3/5 vs 2/5 feature、4/5 model refuted、5/5 data inconclusive、rollback regret 0.2586、v5 FAIL、Caco2 pending confirmation", "不要写五轴互补性已由统一消融证明；LOAO 尚未完成"],
    ["R4", "‘detects scientific failure modes’ 过宽；证据主要是本地自编 fixture", "ADMET-EvO detects prespecified failures and prevents unsafe data use", "主文聚焦 hidden faults、data stewardship、rollback；committee replay 放补充；UQ 可移 R5", "rules parity、unsafe merge=0、composite challenge 的 local/sealed 限定", "没有独立 expert pack 前不得写 broad/independent/general scientific failure detection"],
    ["R5", "‘unknown endpoint’ 包装不准确；真正未知的是 blind labels/molecules", "A prospective multi-endpoint CYP study freezes scientific decisions before external evaluation", "按 data contract→identity split→Stage A/B→external data route→internal certify→blind freeze 排序", "数据规模、B1/M3/M4、6 个 external-data deltas、4/4 internal gates、750-row hashes", "不能写 unknown endpoint；不能写 external blind performance/rank 已经证明"],
]


GAP_HEADERS = ["优先级", "Result", "缺口", "为什么关键", "最小交付", "完成后可升级的主张", "当前状态"]
GAP_ROWS = [
    ["P0", "R5", "OpenADMET external blind readout", "决定 prospective case study 是否有外部有效性", "官方 score/rank、每 endpoint 指标、CV→blind gap、一次性锁定分析", "从 frozen workflow 升级到 externally evaluated prospective study", "未完成"],
    ["P0", "R4", "Independent expert blind fault pack", "当前 fault/data challenge 为本地设计，与 complete rules 持平", "≥12 unseen cases；4 clean + 8 fault；双人 blind scoring", "从 prespecified/local failures 升级到 independent failure handling", "协议/author kit 已有，结果未返回"],
    ["P0", "R3", "Five-axis leave-one-axis-out", "当前五轴证据来自分散实验，无法支持系统级互补性", "Full、−Feature、−Model、−Data、−Replicate、−Stop；matched budget/cost", "五类 research actions provide complementary system-level capabilities", "未执行"],
    ["P0", "R3", "Caco2 external-data source lineage clearance + remaining v2 confirmations", "最新 22-task data action 目前主要是 development-only", "清理 Caco2 lineage；在独立 seeds/冻结最优 recipe 上确认预注册任务", "安全地加入跨 endpoint adaptive external-data 结果", "Caco2 数值完成但 lineage pending；其余未完成"],
    ["P0", "R2", "统一主性能表替换旧网页排名叙述", "网页存在 Bioavailability 异常，且 live leaderboard 容易变化", "主文采用 8-model source table；snapshot 作为补充", "broad competitiveness（来源可追溯）", "数据已完成，需改稿"],
    ["P1", "R1", "Hardness multi-seed stability", "当前 hardness 是单一 discovery profile，稳定性未知", "22 endpoints × 5 profiling seeds；ICC/Spearman/Fleiss kappa", "stable/reproducible hardness diagnosis", "未完成"],
    ["P1", "R2", "端到端资源成本", "72.2% 只覆盖 selected-recipe fitting 且由 LD50 主导", "CPU/GPU hours、peak memory、LLM tokens/API cost、wall-clock；失败成本入账", "end-to-end compute efficiency", "部分字段缺失"],
    ["P1", "R2/R3", "Agent policy vs AutoML 的新独立确认", "现有正式 noninferiority 与多个 promotion gate 失败", "仅在新方案预注册 gate 通过后解封独立 cohort", "policy predictive value", "当前不可宣称；不建议为补正结果追跑"],
    ["P2", "R4/R5", "统一 applicability/abstention policy", "combined UQ 在 0/6 endpoint 未胜单信号", "endpoint-conditional signal selection + coverage-risk curve", "validated abstention/UQ policy", "内部审计完成，promotion 失败"],
]


SECTION_TABLE_HEADERS = [
    "Result",
    "建议小章节",
    "正文数据表目的",
    "建议行",
    "建议列",
    "现有数据支撑",
    "仍需补的数据",
    "推荐位置",
]
SECTION_TABLE_ROWS = [
    ["R1", "1.1 Endpoint definition and hardness diagnosis", "把框架描述转成可量化系统结果", "22 endpoints；可在主文聚合、补充材料逐 endpoint 展开", "endpoint/category/task type/metric/contract completeness/hardness/domain shift/imbalance/roughness/routed action", "22/22 contracts；hardness mean 0.313；17/4/1 routing；domain shift 0.748", "5-seed hardness ICC、Spearman、action kappa", "主文 Table 1 + Supplementary task-level table"],
    ["R1", "1.2 Auditable lineage and immutable certification", "证明失败留存、篡改防护和 freeze 边界", "ledger/audit/freeze 三类证据", "artifact/count/pass condition/result/hash/test_loaded", "834 records；789 completed；1 failure；570 passed/1 skipped；22/22 freeze", "若代码继续变化，需重新跑冻结测试套件", "主文小表或 Figure 1 audit inset"],
    ["R2", "2.1 Broad predictive competitiveness", "给出统一、来源可追溯的 22-task 性能比较", "22 tasks + 8-model summary row", "task/metric/ADMET-EvO/7 comparators/within-table rank/best", "8-model 完整表已完成；ADMET-EvO rank 2/8、mean 2.977、best 4", "无；captured-board rank 仅作 sensitivity", "主文核心性能表"],
    ["R2", "2.2 Targeted expert acquisition", "显示模型动作的 endpoint 条件性", "Bioavailability、CYP2C9 substrate、Lipophilicity 三行", "endpoint/action/control/n/mean delta/95% CI/wins/runtime ratio/interpretation", "GNN 两个 paired contrasts；UniMol2 164M vs 84M", "最好增加独立 endpoint replicate；否则保持 post-hoc 限定", "主文小表或 Extended Data"],
    ["R2", "2.3 Compute-aware selection and policy controls", "同时展示效率收益与 policy 负结果", "efficiency、Agent-vs-AutoML、memory、random 四行", "contrast/n/mean delta/95% CI/W-T-L/cost ratio/gate", "22-task efficiency与三项 control 已完成", "端到端 GPU/CPU/LLM/wall-clock 成本", "主文核心机制表；负结果不可省略"],
    ["R3", "3.1 Complete action-effect grid", "防止只展示 action 正例", "Feature/Model/Data 三行", "axis/endpoints/seeds/supported/inconclusive/refuted/posterior/cost", "100 fits；Feature 3/0/2；Model 0/1/4；Data 0/5/0", "无；完整 task-level effects 可放 Supplement", "主文机制表"],
    ["R3", "3.2 Policy falsification and rollback", "把 self-evolution 落到可观测策略变化", "prior、fresh confirmation、rollback 三类行", "prior/action/fresh endpoints/positive regret/max regret/gate/policy change", "50 fits；2 supported；max regret 0.258620；rollback", "若声称 sustained improvement，仍需后续独立 cohort outcome", "主文自我修正小表/流程图"],
    ["R3", "3.3 Adaptive external-data actions", "区分开发筛选、独立确认和未完成任务", "22 task screening + confirmed tasks", "task/source/decision/ratio/development delta/confirmation delta/positive seeds/lineage/status", "v1：9 augment/13 abstain；v2 Caco2 +0.010208（3/4 positive）", "Caco2 lineage clearance；其余冻结最优 recipe 独立 confirmation", "Supplementary 大表；确认后精选主文 1–3 个案例"],
    ["R4", "4.1 Hidden fault detection", "对比 Agent、rules、AutoML 与 ablation", "10 cases 或按 fault family 聚合", "fault/severity/agent/rules/schema-AutoML/AutoML/detection/root cause/remediation", "Agent F1=1；rules=1；expanded AutoML 0.615；no-hardness=0", "独立 expert blind pack；clean false-alarm", "主文安全表；本地 case 明确标注"],
    ["R4", "4.2 Data stewardship", "展示安全合并与 reject-all 的机会损失", "11 dossiers + 方法汇总", "decision accuracy/issue F1/manifest accuracy/unsafe merges/useful accepts", "Agent 全 1.0/unsafe 0；schema unsafe 7；direct-fit unsafe 10", "外部独立 dossiers 更佳但非必需", "主文安全表或与 4.1 合并"],
    ["R4", "4.3 Bounded correction and stopping", "显示修复、回滚、停止不是事后叙述", "self-correction 2 cases + post-memory 5 endpoints + composite 10 cases", "trigger/action/child delta/stop/compute/unsafe/rollback", "2/2 repair；12 fits；两条 stop；composite macro 1.0/unsafe 0", "外部 repair outcome", "主文小表；committee replay 放 Supplement"],
    ["R5", "5.1 CYP data contracts and leakage-safe split", "建立 prospective case study 的样本与评估基础", "6 endpoints/modalities + blind input", "endpoint/assay/label/task type/rows/valid labels/positive rate/identities/groups/folds", "4,905/6,145/6,146/17,504/750；6,145 identities；3,869 groups", "无", "主文 case-study data table"],
    ["R5", "5.2 Model and external-data decisions", "显示 agent 否证复杂模型并按任务类型路由数据", "B1/M3/M4 + 6 endpoint external treatment rows", "candidate/metric/CV result/gate；endpoint/task type/treatment delta/router decision", "Stage B 与 4 正/2 负 external-data deltas 已完成", "blind readout 返回前不得添加 test-driven selection", "主文决策表"],
    ["R5", "5.3 Internal certification and blind freeze", "区分内部通过与外部结果待定", "4 internal gates + 2 submission files", "gate/threshold/observed/pass；file/rows/columns/SHA/schema/feedback", "4/4 PASS；两份 750-row 文件 schema PASS/hash frozen", "官方 blind score/rank、CV→blind gap", "主文末表；外部 readout 到位后追加一列"],
]


README_ROWS = [
    ["文件用途", "ADMET-EvO 五个 Results 的证据主张、已统计数据、来源、限制和实验缺口总表。"],
    ["证据截点", "2026-08-25 UTC；包含 2026-08-25 新增 8-model 对照、计算效率 Excel、adaptive-data v1/v2 当前状态。"],
    ["主文可用", "证据已形成可核验 artifact；仍需按‘限制’列约束措辞。"],
    ["主文可用（需限定）", "可进入主文，但必须写明 validation/internal/local/snapshot 等证据边界。"],
    ["负结果（必须披露）", "实验已完成但 gate 失败；它是科研结果，不应删除或包装成成功。"],
    ["探索性/待独立确认", "仅开发期或筛选期结果；不得升级为 confirmatory claim。"],
    ["未完成/不可宣称", "只有协议、数据冻结或部分任务；不能写成已经完成的结果。"],
    ["跨任务汇总原则", "不同 endpoint 的 AUROC/AUPRC/Spearman/MAE 原始 delta 不直接求平均；先定向并按预注册尺度归一化。"],
    ["排序原则", "主文优先使用来源可追溯、预注册、独立或冻结证据；历史网页排名和开发期筛选降级为补充材料。"],
]


STATUS_COLORS = {
    "主文可用": "C6EFCE",
    "主文可用（需限定）": "FFF2CC",
    "主文可用（需弱化）": "FFF2CC",
    "主文可用（负结果）": "FCE4D6",
    "负结果（必须披露）": "F4B183",
    "补充材料": "D9EAF7",
    "补充材料/需限定": "D9EAF7",
    "探索性/待独立确认": "DDEBF7",
    "待来源许可/lineage 清理": "FFE699",
    "未完成/不可宣称": "F4CCCC",
}


def add_table(ws, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def style_sheet(ws, widths=None, status_col=None, freeze="A2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E1F2")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    if status_col:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=status_col)
            color = STATUS_COLORS.get(str(cell.value), None)
            if color:
                cell.fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    if widths:
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 34


def append_rows(ws, headers, rows) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)


def import_csv_sheet(wb, title: str, csv_path: Path, table_name: str) -> None:
    ws = wb.create_sheet(title)
    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.reader(handle):
            converted = []
            for value in row:
                try:
                    converted.append(float(value) if value != "" else None)
                except ValueError:
                    converted.append(value)
            ws.append(converted)
    widths = [28] + [16] * max(1, ws.max_column - 1)
    style_sheet(ws, widths=widths)
    add_table(ws, table_name)


def import_existing_xlsx_sheet(wb, title: str, xlsx_path: Path, source_sheet: str) -> None:
    src_wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    src_ws = src_wb[source_sheet]
    ws = wb.create_sheet(title)
    rows = list(src_ws.iter_rows(values_only=True))
    # The source has a title row followed by the actual table header.
    for row in rows[1:]:
        ws.append(list(row))
    style_sheet(ws, widths=[32, 14, 18, 18, 18, 20, 20, 14, 16, 22, 34, 34, 20, 20])
    add_table(ws, "Compute22Table")


def build_workbook() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results证据主表"
    append_rows(ws, MASTER_HEADERS, MASTER_ROWS)
    style_sheet(
        ws,
        widths=[8, 12, 34, 38, 28, 35, 26, 54, 24, 22, 46, 48, 62, 10],
        status_col=10,
    )
    add_table(ws, "ResultsEvidenceMaster")

    ws_rewrite = wb.create_sheet("Results改写建议")
    append_rows(ws_rewrite, REWRITE_HEADERS, REWRITE_ROWS)
    style_sheet(ws_rewrite, widths=[10, 38, 52, 58, 58, 58])
    add_table(ws_rewrite, "ResultsRewriteTable")

    ws_gap = wb.create_sheet("缺口优先级")
    append_rows(ws_gap, GAP_HEADERS, GAP_ROWS)
    style_sheet(ws_gap, widths=[10, 10, 38, 52, 58, 52, 32])
    add_table(ws_gap, "ResultsGapTable")

    ws_sections = wb.create_sheet("小章节数据支撑表")
    append_rows(ws_sections, SECTION_TABLE_HEADERS, SECTION_TABLE_ROWS)
    style_sheet(ws_sections, widths=[10, 38, 46, 52, 62, 58, 52, 34])
    add_table(ws_sections, "SectionSupportTable")

    import_csv_sheet(
        wb,
        "TDC22_8模型",
        REPORTS / "model_comparison_tdc22_8models_2026-08-25.csv",
        "TDC22EightModels",
    )
    import_existing_xlsx_sheet(
        wb,
        "Compute22明细",
        REPORTS / "ADMET_EvO_compute_aware_selection_summary.xlsx",
        "22任务明细",
    )
    import_csv_sheet(
        wb,
        "AdaptiveData22",
        REPORTS / "all22_adaptive_agent_data_v1" / "all22_data_augmentation_summary.csv",
        "AdaptiveData22Table",
    )

    ws_readme = wb.create_sheet("使用说明", 0)
    append_rows(ws_readme, ["项目", "说明"], README_ROWS)
    style_sheet(ws_readme, widths=[30, 110])
    add_table(ws_readme, "WorkbookReadme")

    for ws_item in wb.worksheets:
        ws_item.sheet_properties.pageSetUpPr.fitToPage = True
        ws_item.page_setup.fitToWidth = 1
        ws_item.page_setup.fitToHeight = 0
        ws_item.sheet_properties.outlinePr.summaryBelow = True
    wb.save(OUT_XLSX)


def write_master_csv() -> None:
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(MASTER_HEADERS)
        writer.writerows(MASTER_ROWS)


if __name__ == "__main__":
    write_master_csv()
    build_workbook()
    print(OUT_XLSX)
    print(OUT_CSV)
