"""Build the v2 audit workbook for optimal-recipe data augmentation."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
MAPPING = ROOT / "runs/v27-maplight-gnn-freeze/full-22-task.mapping.json"
SCORECARD = ROOT / "reports/goal_scorecard_v27.json"
V2_ROOT = ROOT / "reports/all22_optimal_recipe_adaptive_data_v2"
OUTPUT = V2_ROOT / "all22_optimal_recipe_augmentation_audit.xlsx"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _recipe_label(frozen: dict) -> str:
    if frozen.get("kind") == "recipe":
        return f"recipe({frozen['recipe']})"
    if frozen.get("kind") == "ensemble":
        return "ensemble(" + "+".join(row["model"] for row in frozen["members"]) + ")"
    return str(frozen["model"])


def _runtime_status(frozen: dict) -> str | None:
    if frozen.get("kind") != "recipe":
        return None
    recipe = str(frozen["recipe"])
    if "finetuned_unimol2" in recipe:
        return "blocked: CUDA unavailable for frozen Uni-Mol2 fine-tuning"
    if "maplight_gnn" in recipe:
        return "blocked: frozen MapLight+GNN isolated runtime unavailable"
    return None


def _artifact(directory: str, task: str) -> dict | None:
    path = V2_ROOT / directory / f"{task}.json"
    return _load_json(path) if path.exists() else None


def _best_observed(development: dict | None) -> tuple[float | None, float | None]:
    if not development or not development.get("ratio_summaries"):
        return None, None
    ratio, summary = max(
        development["ratio_summaries"].items(),
        key=lambda item: item[1]["oriented_delta_mean"],
    )
    return float(ratio), float(summary["oriented_delta_mean"])


def main() -> None:
    mapping = _load_json(MAPPING)
    scorecard_rows = _load_json(SCORECARD)["tasks"]
    scorecard = {row["task"]: row for row in scorecard_rows}

    headers = [
        "Task",
        "Metric",
        "Frozen v2.7 Recipe",
        "v2.7 Official Score (reference only)",
        "Corrected Dev Baseline",
        "Corrected Selected Ratio",
        "Corrected Dev Augmented",
        "Corrected Dev Oriented Δ",
        "Corrected Best Observed Ratio",
        "Corrected Best Observed Δ",
        "Development Decision",
        "Confirmation Baseline",
        "Confirmation Augmented",
        "Confirmation Oriented Δ",
        "Confirmation Positive Seeds",
        "Corrected Final Status",
    ]
    output_rows = []
    for score_row in scorecard_rows:
        task = score_row["task"]
        development = _artifact("development", task)
        confirmation = _artifact("confirmation", task)
        selected_ratio = None
        dev_baseline = None
        dev_augmented = None
        dev_delta = None
        dev_decision = "not_run"
        best_ratio, best_delta = _best_observed(development)
        if development:
            dev_decision = development["status"]
            summaries = development.get("ratio_summaries", {})
            if summaries:
                dev_baseline = next(iter(summaries.values()))["baseline_mean"]
            selected_ratio = development.get("selection", {}).get("selected_ratio")
            if selected_ratio is not None:
                selected = summaries[str(float(selected_ratio))]
                dev_augmented = selected["augmented_mean"]
                dev_delta = selected["oriented_delta_mean"]
        blocker = _runtime_status(mapping[task]) if development is None else None
        if confirmation:
            summary = confirmation["agent_summary"]
            final_status = (
                "retain_augmented_confirmed"
                if confirmation["status"] == "promoted"
                else "retain_baseline_confirmation_failed"
            )
        else:
            summary = {}
            if development and selected_ratio is None:
                final_status = "retain_baseline_development_abstained"
            elif development:
                final_status = "confirmation_pending"
            elif blocker:
                final_status = blocker
            else:
                final_status = "development_pending"
        output_rows.append(
            [
                task,
                score_row["metric"],
                _recipe_label(mapping[task]),
                float(scorecard[task]["score"]),
                dev_baseline,
                selected_ratio,
                dev_augmented,
                dev_delta,
                best_ratio,
                best_delta,
                dev_decision,
                summary.get("baseline_mean"),
                summary.get("augmented_mean"),
                summary.get("oriented_delta_mean"),
                summary.get("positive_seed_fraction"),
                final_status,
            ]
        )

    workbook = Workbook()
    results = workbook.active
    results.title = "Corrected 22-task audit"
    results.append(headers)
    for row in output_rows:
        results.append(row)
    results.freeze_panes = "A2"
    results.auto_filter.ref = results.dimensions
    results.row_dimensions[1].height = 42
    for cell in results[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [34, 11, 52, 22, 22, 20, 22, 22, 22, 21, 33, 22, 22, 23, 22, 45]
    for index, width in enumerate(widths, start=1):
        results.column_dimensions[results.cell(1, index).column_letter].width = width
    for row in results.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in (4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15):
            row[index - 1].number_format = "0.000000"
    end_row = results.max_row
    table = Table(displayName="OptimalRecipeAugmentationAudit", ref=f"A1:P{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    results.add_table(table)
    for column in ("H", "J", "N"):
        results.conditional_formatting.add(
            f"{column}2:{column}{end_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="num",
                mid_value=0,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    notes = workbook.create_sheet("Protocol notes")
    notes_rows = [
        ("Item", "Value"),
        ("Why v1 is invalid as the final Agent comparison", "v1 used one Morgan radius-2 + LightGBM controller on all 22 tasks, not each task's frozen optimal v2.7 recipe."),
        ("Correct comparison", "Within each task and seed, native-only and augmented arms use the identical frozen v2.7 recipe, representation, hyperparameters, split and seed; only admitted training membership changes."),
        ("Development", "TDC default scaffold seeds 2-5; ratios 0.5x, 1x, 2x, 4x, 8x; practical gate oriented delta >=0.005 and >=75% positive seeds."),
        ("Confirmation", "Disjoint seeds 8-11; selected ratio only."),
        ("Official score caveat", "The v2.7 official score is shown only to diagnose the weak fixed baseline. It is an official-test result and must not be subtracted from scaffold-validation scores as a paired effect."),
        ("Current runtime limitation", "No allocated CUDA device; frozen MapLight+GNN isolated environment is absent. Deep-recipe tasks remain blocked instead of being silently replaced by weaker models."),
        ("Source of truth", str(MAPPING.relative_to(ROOT))),
        ("Development artifacts", str((V2_ROOT / "development").relative_to(ROOT))),
        ("Confirmation artifacts", str((V2_ROOT / "confirmation").relative_to(ROOT))),
    ]
    for row in notes_rows:
        notes.append(row)
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 120
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
