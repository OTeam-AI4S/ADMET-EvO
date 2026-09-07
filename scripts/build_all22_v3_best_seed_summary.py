#!/usr/bin/env python3
"""Build the requested optimistic best-seed/ratio summary for the fixed v3 grid."""

from __future__ import annotations

import csv
import json
import math
import statistics
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from run_all22_adaptive_agent_data import TASKS
from run_all22_optimal_recipe_adaptive_data import PROTOCOL


ROOT = Path(__file__).resolve().parents[1]
REPORT_ROOT = ROOT / "reports/all22_optimal_recipe_adaptive_data_v3"
DEVELOPMENT = REPORT_ROOT / "development"
CSV_OUTPUT = REPORT_ROOT / "all22_best_seed_ratio_summary.csv"
XLSX_OUTPUT = REPORT_ROOT / "all22_best_seed_ratio_summary.xlsx"
AUDIT_OUTPUT = REPORT_ROOT / "all22_best_seed_ratio_summary.audit.json"
FIRST_VARIANCE_SEEDS = (2, 3, 4)

HEADERS = [
    "Task",
    "Metric",
    "Frozen Recipe",
    "Development Status",
    "Best Seed",
    "Best Ratio",
    "Baseline Score",
    "Augmented Score",
    "Oriented Delta",
    "Variance Seeds",
    "Augmented Score Sample Variance",
    "Augmented Score Sample SD",
    "Oriented Delta Sample Variance",
    "Oriented Delta Sample SD",
    "External-Test Exact Rows Excluded",
    "Post-Firewall External Identities",
    "Test Structure Columns Loaded",
    "Test Labels Loaded",
    "Development Artifact",
]


def _load_task(task: str) -> tuple[Path, dict[str, Any]]:
    path = DEVELOPMENT / f"{task}.json"
    error = DEVELOPMENT / f"{task}.error.json"
    if error.exists():
        raise RuntimeError(f"{task} has an execution error artifact: {error}")
    if not path.exists():
        raise FileNotFoundError(f"Missing development artifact: {path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    safety = payload.get("safety", {})
    if not (
        payload.get("protocol") == PROTOCOL
        and safety.get("tdc_test_loaded") is True
        and safety.get("tdc_test_structures_loaded_for_firewall") is True
        and safety.get("tdc_test_columns_loaded") == ["Drug"]
        and safety.get("tdc_test_labels_loaded") is False
        and safety.get("test_labels_accessible") is False
    ):
        raise RuntimeError(f"{task} failed the v3 structure-only firewall audit")
    return path, payload


def _sample_variance(values: list[float]) -> tuple[float | None, float | None]:
    if len(values) < 2:
        return None, None
    variance = float(statistics.variance(values))
    return variance, float(math.sqrt(variance))


def _summary_row(task: str) -> tuple[list[Any], dict[str, Any]]:
    path, payload = _load_task(task)
    rows = [
        row
        for row in payload.get("per_seed_ratio_model", [])
        if math.isfinite(float(row["oriented_delta"]))
    ]
    curation = payload.get("curation", {})
    safety = payload["safety"]
    if not rows:
        output = [
            task,
            payload.get("metric"),
            payload.get("frozen_recipe_label"),
            payload.get("status"),
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            curation.get("exact_tdc_test_structure_overlap_exclusions"),
            curation.get("post_exact_firewall_identities"),
            ",".join(safety["tdc_test_columns_loaded"]),
            safety["tdc_test_labels_loaded"],
            str(path.resolve()),
        ]
        return output, {"task": task, "status": payload.get("status"), "best": None}

    best = max(
        rows,
        key=lambda row: (
            float(row["oriented_delta"]),
            -float(row["ratio"]),
            -int(row["seed"]),
        ),
    )
    ratio = float(best["ratio"])
    variance_rows = [
        row
        for row in rows
        if int(row["seed"]) in FIRST_VARIANCE_SEEDS
        and math.isclose(float(row["ratio"]), ratio, rel_tol=0.0, abs_tol=1e-12)
    ]
    variance_rows.sort(key=lambda row: int(row["seed"]))
    if [int(row["seed"]) for row in variance_rows] != list(FIRST_VARIANCE_SEEDS):
        raise RuntimeError(f"{task} lacks first-three-seed rows at ratio {ratio:g}")
    score_variance, score_sd = _sample_variance(
        [float(row["augmented_score"]) for row in variance_rows]
    )
    delta_variance, delta_sd = _sample_variance(
        [float(row["oriented_delta"]) for row in variance_rows]
    )
    output = [
        task,
        payload.get("metric"),
        payload.get("frozen_recipe_label"),
        payload.get("status"),
        int(best["seed"]),
        ratio,
        float(best["baseline_score"]),
        float(best["augmented_score"]),
        float(best["oriented_delta"]),
        ",".join(str(seed) for seed in FIRST_VARIANCE_SEEDS),
        score_variance,
        score_sd,
        delta_variance,
        delta_sd,
        curation.get("exact_tdc_test_structure_overlap_exclusions"),
        curation.get("post_exact_firewall_identities"),
        ",".join(safety["tdc_test_columns_loaded"]),
        safety["tdc_test_labels_loaded"],
        str(path.resolve()),
    ]
    audit = {
        "task": task,
        "status": payload.get("status"),
        "best": {
            "seed": int(best["seed"]),
            "ratio": ratio,
            "baseline_score": float(best["baseline_score"]),
            "augmented_score": float(best["augmented_score"]),
            "oriented_delta": float(best["oriented_delta"]),
        },
        "variance_seeds": list(FIRST_VARIANCE_SEEDS),
        "augmented_score_sample_variance": score_variance,
        "oriented_delta_sample_variance": delta_variance,
    }
    return output, audit


def main() -> None:
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    rows: list[list[Any]] = []
    audits: list[dict[str, Any]] = []
    for task in TASKS:
        row, audit = _summary_row(task)
        rows.append(row)
        audits.append(audit)

    with CSV_OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Best seed ratio"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [34, 12, 54, 34, 12, 12, 18, 18, 18, 16, 24, 20, 25, 20, 25, 25, 24, 18, 75]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in range(6, 15):
            row[index - 1].number_format = "0.000000"
    table = Table(displayName="V3BestSeedRatio", ref=f"A1:S{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    notes = workbook.create_sheet("Notes")
    notes.append(["Item", "Value"])
    notes.append(["Protocol", PROTOCOL])
    notes.append(["Grid", "Development seeds 2,3,4,5 × ratios 0.5,1,2,4,8"])
    notes.append([
        "Selection",
        "Per user request: maximize the single observed oriented delta; ties use the smaller ratio, then smaller seed.",
    ])
    notes.append([
        "Variance",
        "Sample variance and SD across seeds 2,3,4 at the selected best ratio.",
    ])
    notes.append([
        "Interpretation",
        "This is an optimistic post-hoc best-seed summary, not an unbiased multi-seed estimate.",
    ])
    notes.append([
        "Leakage firewall",
        "Only hidden-test Drug structures are loaded for exact identity exclusion; test labels are not loaded.",
    ])
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 120
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(XLSX_OUTPUT)

    audit_payload = {
        "protocol": PROTOCOL,
        "selection_rule": "maximum single observed oriented delta",
        "tie_break": "smaller ratio, then smaller seed",
        "variance_seeds": list(FIRST_VARIANCE_SEEDS),
        "task_count": len(audits),
        "tasks": audits,
        "csv": str(CSV_OUTPUT.resolve()),
        "xlsx": str(XLSX_OUTPUT.resolve()),
    }
    AUDIT_OUTPUT.write_text(
        json.dumps(audit_payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(audit_payload, sort_keys=True))


if __name__ == "__main__":
    main()
