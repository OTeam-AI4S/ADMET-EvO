"""Build the complete ADMET322 holdout workbook with numerical charts."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY, BLUE, TEAL, GREEN, GOLD, RED = "17365D", "2F75B5", "2F8F9D", "70AD47", "FFC000", "C00000"
WHITE, LIGHT_BLUE, LIGHT_GREEN, LIGHT_GOLD, LIGHT_RED = "FFFFFF", "D9EAF7", "E2F0D9", "FFF2CC", "FCE4D6"
GRID = Side(style="thin", color="D9E1F2")
MODELS = (
    "extra_trees_morgan", "extra_trees_rich", "lightgbm_rich",
    "random_forest_rich", "xgboost_rich", "catboost_rich",
)


def _append(ws, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))


def _style(ws) -> None:
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=GRID)
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 18
    ws.column_dimensions["A"].width = 54


def build(summary_path: Path, development_path: Path, output: Path) -> None:
    holdout = pd.read_csv(summary_path)
    development = pd.read_csv(development_path)
    if len(holdout) != 322 or len(development) != 372:
        raise RuntimeError("Workbook requires the complete 322-task holdout and 372-task development denominators")
    holdout = holdout.sort_values(
        ["agent_mcc", "agent_gain_vs_morgan"], ascending=[False, False], kind="stable"
    )
    tier_order = ["strong", "good", "positive", "no_gain"]
    tiers = holdout["holdout_tier"].value_counts().reindex(tier_order, fill_value=0).rename_axis("Tier").reset_index(name="Endpoints")
    wtl_rows = []
    for model in MODELS:
        delta = holdout["agent_mcc"] - holdout[f"{model}_mcc"]
        wtl_rows.append({
            "Comparator": model,
            "Agent wins": int(delta.gt(1e-12).sum()),
            "Ties": int(delta.abs().le(1e-12).sum()),
            "Agent losses": int(delta.lt(-1e-12).sum()),
            "Mean MCC delta": float(delta.mean()),
            "Median MCC delta": float(delta.median()),
        })
    wtl = pd.DataFrame(wtl_rows)
    family = holdout.groupby(["source", "assay_family"]).agg(
        endpoints=("endpoint_id", "size"),
        strong=("holdout_tier", lambda x: int(x.eq("strong").sum())),
        good_or_strong=("holdout_tier", lambda x: int(x.isin(["strong", "good"]).sum())),
        positive_gain=("agent_gain_vs_morgan", lambda x: int(x.gt(0).sum())),
        fusion_selected=("selected_kind", lambda x: int(x.eq("fusion").sum())),
        mean_agent_mcc=("agent_mcc", "mean"),
        mean_gain_vs_morgan=("agent_gain_vs_morgan", "mean"),
    ).reset_index()
    family["strong_rate"] = family["strong"] / family["endpoints"]
    family = family.sort_values(["strong_rate", "mean_agent_mcc"], ascending=False)
    top20 = holdout.loc[holdout["holdout_tier"].eq("strong")].head(20).copy()

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "ADMET322 Fresh Endpoint Fusion Confirmation"
    workbook.properties.creator = "ADMET Agent"
    overview = workbook.create_sheet("Overview")
    overview.merge_cells("A1:H1")
    overview["A1"] = "322个全新ADMET毒理学任务：冻结Agent配方的严格holdout确认"
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].font = Font(color=WHITE, bold=True, size=16)
    overview["A1"].alignment = Alignment(horizontal="center")
    overview.append([])
    overview.append(["Protocol item", "Value"])
    overview.append(["Fresh denominator", "322 endpoints: 309 ToxCast + 13 SIDER; all remain visible"])
    overview.append(["Development", "Folds 0–2 only, seed 17; 6 experts + 45 fixed pair blends per endpoint"])
    overview.append(["Confirmation", "Frozen recipe fit on folds 0–2; untouched folds 3–4; seeds 29 and 43"])
    overview.append(["Strong gate", "Mean MCC ≥ 0.25; gain vs Morgan ExtraTrees ≥ 0.05; wins on both confirmation seeds"])
    overview.append(["Strong endpoints", int(holdout["holdout_tier"].eq("strong").sum())])
    overview.append(["Good or strong", int(holdout["holdout_tier"].isin(["strong", "good"]).sum())])
    overview.append(["Positive gain", int(holdout["agent_gain_vs_morgan"].gt(0).sum())])
    overview.append(["Fusion selected", int(holdout["selected_kind"].eq("fusion").sum())])
    overview.append(["Mean Agent MCC", float(holdout["agent_mcc"].mean())])
    overview.append(["Mean Morgan ET MCC", float(holdout["morgan_et_mcc"].mean())])
    overview.append(["Top20 warning", "Top20 is a post-hoc presentation view only; primary evidence is the complete 322-task denominator"])
    for cell in overview[3]:
        cell.fill = PatternFill("solid", fgColor=TEAL); cell.font = Font(color=WHITE, bold=True)
    for row in overview.iter_rows(min_row=4):
        row[0].font = Font(bold=True, color=NAVY); row[1].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    overview.column_dimensions["A"].width = 30; overview.column_dimensions["B"].width = 110

    full = workbook.create_sheet("All 322 Holdout")
    _append(full, holdout); _style(full)
    columns = {cell.value: cell.column for cell in full[1]}
    for row in range(2, full.max_row + 1):
        tier = full.cell(row, columns["holdout_tier"])
        tier.fill = PatternFill("solid", fgColor={"strong": LIGHT_GREEN, "good": LIGHT_BLUE, "positive": LIGHT_GOLD, "no_gain": LIGHT_RED}[tier.value])
    for name in ["agent_mcc", "agent_gain_vs_morgan"]:
        col = get_column_letter(columns[name])
        full.conditional_formatting.add(
            f"{col}2:{col}{full.max_row}",
            ColorScaleRule(start_type="min", start_color=RED, mid_type="percentile", mid_value=50,
                           mid_color=LIGHT_GOLD, end_type="max", end_color=GREEN),
        )

    strong = workbook.create_sheet("Strong Tasks")
    _append(strong, holdout.loc[holdout["holdout_tier"].eq("strong")]); _style(strong)
    top = workbook.create_sheet("Top20 Presentation")
    _append(top, top20); _style(top)
    dev = workbook.create_sheet("All 372 Development")
    _append(dev, development); _style(dev)
    fam = workbook.create_sheet("Assay Families")
    _append(fam, family); _style(fam)
    wtl_sheet = workbook.create_sheet("Model Win-Tie-Loss")
    _append(wtl_sheet, wtl); _style(wtl_sheet)

    chart_data = workbook.create_sheet("Chart Data")
    _append(chart_data, tiers)
    family_start = 5
    for offset, name in enumerate(family.columns):
        chart_data.cell(1, family_start + offset, name)
    for row_index, row in enumerate(family.itertuples(index=False, name=None), start=2):
        for offset, value in enumerate(row):
            chart_data.cell(row_index, family_start + offset, value)
    top_start = family_start + len(family.columns) + 2
    top_chart = top20[["endpoint_id", "agent_mcc", "morgan_et_mcc"]]
    for offset, name in enumerate(top_chart.columns):
        chart_data.cell(1, top_start + offset, name)
    for row_index, row in enumerate(top_chart.itertuples(index=False, name=None), start=2):
        for offset, value in enumerate(row):
            chart_data.cell(row_index, top_start + offset, value)
    wtl_start = top_start + len(top_chart.columns) + 2
    for offset, name in enumerate(wtl.columns):
        chart_data.cell(1, wtl_start + offset, name)
    for row_index, row in enumerate(wtl.itertuples(index=False, name=None), start=2):
        for offset, value in enumerate(row):
            chart_data.cell(row_index, wtl_start + offset, value)
    _style(chart_data)

    charts = workbook.create_sheet("Numeric Charts")
    charts["A1"] = "ADMET322 strict holdout numerical comparisons"
    charts["A1"].font = Font(size=16, bold=True, color=NAVY)
    tier_chart = PieChart(); tier_chart.title = "Complete 322-task success tiers"
    tier_chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=1 + len(tiers)), titles_from_data=True)
    tier_chart.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=1 + len(tiers)))
    tier_chart.dLbls = DataLabelList(); tier_chart.dLbls.showVal = True; tier_chart.dLbls.showPercent = True
    tier_chart.height = 9; tier_chart.width = 14; charts.add_chart(tier_chart, "A3")
    score_chart = BarChart(); score_chart.type = "bar"; score_chart.grouping = "clustered"
    score_chart.title = "Top20 presentation: Agent vs Morgan ExtraTrees"
    score_chart.x_axis.title = "Holdout MCC"; score_chart.height = 15; score_chart.width = 23
    score_chart.add_data(Reference(chart_data, min_col=top_start + 1, max_col=top_start + 2, min_row=1, max_row=21), titles_from_data=True)
    score_chart.set_categories(Reference(chart_data, min_col=top_start, min_row=2, max_row=21))
    charts.add_chart(score_chart, "P3")
    wtl_chart = BarChart(); wtl_chart.type = "bar"; wtl_chart.grouping = "stacked"; wtl_chart.overlap = 100
    wtl_chart.title = "Agent vs matched classic models on all 322 endpoints"
    wtl_chart.x_axis.title = "Endpoint count"; wtl_chart.height = 10; wtl_chart.width = 18
    wtl_chart.add_data(Reference(chart_data, min_col=wtl_start + 1, max_col=wtl_start + 3, min_row=1, max_row=1 + len(wtl)), titles_from_data=True)
    wtl_chart.set_categories(Reference(chart_data, min_col=wtl_start, min_row=2, max_row=1 + len(wtl)))
    charts.add_chart(wtl_chart, "A23")
    fam_chart = BarChart(); fam_chart.type = "col"
    fam_chart.title = "Strong success rate by source/assay family"; fam_chart.y_axis.title = "Strong rate"
    strong_rate_col = family_start + list(family.columns).index("strong_rate")
    fam_chart.add_data(Reference(chart_data, min_col=strong_rate_col, min_row=1, max_row=1 + len(family)), titles_from_data=True)
    fam_chart.set_categories(Reference(chart_data, min_col=family_start + 1, min_row=2, max_row=1 + len(family)))
    fam_chart.height = 10; fam_chart.width = 19; charts.add_chart(fam_chart, "T23")

    notes = pd.DataFrame([
        ["Primary denominator", "All 322 endpoints outside the prior ADMET50 experiment; no endpoint was removed using holdout performance."],
        ["Leakage control", "Molecular identity and chemistry group occur in exactly one frozen outer fold."],
        ["Agent recipe", "The six experts and all 45 pair weights were scored only by OOF predictions from folds 0–2."],
        ["Fusion gate", "A pair must gain at least 0.005 over its stronger member and win on at least 2/3 development folds."],
        ["Confirmation", "The frozen recipe and all six matched single experts use identical train and holdout rows; seeds 29/43 quantify fit stochasticity."],
        ["Post-hoc oracle", "best_classic_model_posthoc is descriptive only because it is selected after reading holdout results."],
        ["Correlation", "Multiple ToxCast assays from one family are correlated endpoints, not independent external datasets; family summaries are therefore required."],
        ["Top20", "A presentation subset selected after confirmation; it cannot replace the complete 322-task success rates."],
    ], columns=["Item", "Definition / constraint"])
    methods = workbook.create_sheet("Methods & Notes")
    _append(methods, notes); _style(methods); methods.column_dimensions["B"].width = 115

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--development", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.summary.resolve(), args.development.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
