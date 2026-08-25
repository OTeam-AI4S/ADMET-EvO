"""Build a formatted Excel table from the 22-task augmentation summary CSV."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "reports/all22_adaptive_agent_data_v1/all22_data_augmentation_summary.csv"
OUTPUT = ROOT / "reports/all22_adaptive_agent_data_v1/all22_data_augmentation_results.xlsx"

HEADERS = [
    "Task",
    "Metric",
    "Agent Baseline",
    "Agent Augmented",
    "Agent Oriented Δ",
    "Random Forest Δ",
    "Linear SVM Δ",
    "Morgan MLP Δ",
    "Augmentation Decision",
    "Selected External Ratio",
    "Positive Classic Models",
    "Classic Generalized",
    "Task Type",
    "Decision Reason",
    "Final Policy",
]

SOURCE_ORDER = [0, 2, 6, 7, 8, 9, 10, 11, 3, 5, 12, 13, 1, 4, 14]
NUMERIC_COLUMNS = {3, 4, 5, 6, 7, 8, 10, 11}


def main() -> None:
    with INPUT.open(newline="") as handle:
        rows = list(csv.reader(handle))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "22-task results"
    sheet.append(HEADERS)

    for source_row in rows[1:]:
        raw = [source_row[index] for index in SOURCE_ORDER]
        converted: list[object] = []
        for index, value in enumerate(raw, start=1):
            if index in NUMERIC_COLUMNS and value != "":
                converted.append(float(value))
            else:
                converted.append(value)
        sheet.append(converted)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [34, 12, 17, 18, 18, 18, 17, 17, 23, 21, 22, 20, 15, 34, 29]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
        for column in (3, 4, 5, 6, 7, 8, 10, 11):
            row[column - 1].number_format = "0.000000"

    end_row = sheet.max_row
    table = Table(displayName="ADMET22AugmentationResults", ref=f"A1:O{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    sheet.conditional_formatting.add(
        f"E2:H{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="num",
            mid_value=0,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )
    sheet.conditional_formatting.add(
        f"L2:L{end_row}", CellIsRule(operator="equal", formula=['"yes"'], fill=green)
    )
    sheet.conditional_formatting.add(
        f"L2:L{end_row}", CellIsRule(operator="equal", formula=['"no"'], fill=red)
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
