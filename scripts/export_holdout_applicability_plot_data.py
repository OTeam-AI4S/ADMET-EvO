#!/usr/bin/env python3
"""Export the exact source tables used by the hold-out applicability plots."""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "figures" / "holdout_applicability_candidates"
MOLECULE_SOURCE = OUTPUT_DIR / "holdout_applicability_source.csv"
COVERAGES = np.arange(0.25, 1.001, 0.05)
TASK_TYPES = ("Classification", "Regression")


def endpoint_summaries(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    coverage_records: list[dict] = []
    gain_records: list[dict] = []
    for (task, task_type), task_frame in frame.groupby(
        ["task", "task_type"], sort=False
    ):
        ordered = task_frame.sort_values(
            "nearest_train_tanimoto", ascending=False, kind="stable"
        )
        values = ordered["task_normalized_loss"].to_numpy(dtype=float)
        for coverage in COVERAGES:
            count = max(1, int(math.ceil(len(values) * coverage)))
            normalized_error = float(np.mean(values[:count]))
            coverage_records.append(
                {
                    "task": task,
                    "task_type": task_type,
                    "coverage_retained_fraction": float(coverage),
                    "coverage_retained_pct": float(coverage * 100),
                    "n_molecules_retained": count,
                    "normalized_prediction_error_fraction": normalized_error,
                    "normalized_prediction_error_pct": normalized_error * 100,
                }
            )

        trusted = task_frame.loc[
            task_frame["similarity_quartile"] == 4, "task_normalized_loss"
        ]
        reduction = float(100.0 * (1.0 - trusted.mean()))
        gain_records.append(
            {
                "task": task,
                "task_type": task_type,
                "n_molecules_highest_similarity_quartile": int(len(trusted)),
                "mean_normalized_error_highest_similarity_quartile_fraction": float(
                    trusted.mean()
                ),
                "error_reduction_highest_similarity_quartile_pct": reduction,
                "improved": reduction > 0,
            }
        )
    return pd.DataFrame(coverage_records), pd.DataFrame(gain_records)


def bootstrap_median_interval(
    values: np.ndarray, seed: int
) -> tuple[float, float, float]:
    generator = np.random.RandomState(seed)
    indexes = generator.randint(0, len(values), size=(10000, len(values)))
    medians = np.median(values[indexes], axis=1)
    return (
        float(np.median(values)),
        float(np.quantile(medians, 0.025)),
        float(np.quantile(medians, 0.975)),
    )


def coverage_plot_summary(coverage: pd.DataFrame) -> pd.DataFrame:
    records: list[dict] = []
    for type_index, task_type in enumerate(TASK_TYPES):
        for coverage_value in COVERAGES:
            values = coverage.loc[
                (coverage["task_type"] == task_type)
                & np.isclose(
                    coverage["coverage_retained_fraction"], coverage_value
                ),
                "normalized_prediction_error_fraction",
            ].to_numpy(dtype=float)
            seed = 4100 + type_index * 100 + int(coverage_value * 100)
            median, low, high = bootstrap_median_interval(values, seed)
            records.append(
                {
                    "task_type": task_type,
                    "coverage_retained_pct": float(coverage_value * 100),
                    "n_endpoints": int(len(values)),
                    "median_normalized_prediction_error_fraction": median,
                    "bootstrap_95ci_low_fraction": low,
                    "bootstrap_95ci_high_fraction": high,
                    "median_normalized_prediction_error_pct": median * 100,
                    "bootstrap_95ci_low_pct": low * 100,
                    "bootstrap_95ci_high_pct": high * 100,
                    "bootstrap_seed": seed,
                }
            )
    return pd.DataFrame(records)


def gain_plot_summary(gains: pd.DataFrame) -> pd.DataFrame:
    records: list[dict] = []
    for task_type in TASK_TYPES:
        values = gains.loc[
            gains["task_type"] == task_type,
            "error_reduction_highest_similarity_quartile_pct",
        ].to_numpy(dtype=float)
        records.append(
            {
                "task_type": task_type,
                "n_endpoints": int(len(values)),
                "n_endpoints_improved": int(np.sum(values > 0)),
                "median_error_reduction_pct": float(np.median(values)),
                "q25_error_reduction_pct": float(np.quantile(values, 0.25)),
                "q75_error_reduction_pct": float(np.quantile(values, 0.75)),
            }
        )
    return pd.DataFrame(records)


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column_cells in sheet.columns:
            width = min(
                45,
                max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                + 2,
            )
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    workbook.save(path)


def main() -> None:
    if not MOLECULE_SOURCE.is_file():
        raise FileNotFoundError(MOLECULE_SOURCE)

    molecule_data = pd.read_csv(MOLECULE_SOURCE)
    coverage, gains = endpoint_summaries(molecule_data)
    coverage_summary = coverage_plot_summary(coverage)
    gain_summary = gain_plot_summary(gains)

    coverage_path = OUTPUT_DIR / "scheme_1_coverage_reliability_endpoint_data.csv"
    coverage_summary_path = OUTPUT_DIR / "scheme_1_coverage_reliability_plot_data.csv"
    gains_path = OUTPUT_DIR / "scheme_3_trusted_region_gain_plot_data.csv"
    gain_summary_path = OUTPUT_DIR / "scheme_3_trusted_region_gain_summary.csv"
    workbook_path = OUTPUT_DIR / "holdout_applicability_figure_source_data.xlsx"

    coverage.to_csv(coverage_path, index=False)
    coverage_summary.to_csv(coverage_summary_path, index=False)
    gains.to_csv(gains_path, index=False)
    gain_summary.to_csv(gain_summary_path, index=False)

    readme = pd.DataFrame(
        [
            {
                "sheet": "raw_molecule_data",
                "description": "逐分子级基础数据；每条记录含最近训练集 Tanimoto 相似度、相似度四分位和五个冻结种子的平均点损失。",
            },
            {
                "sheet": "coverage_by_endpoint",
                "description": "上图的端点级输入：每个端点在 25%–100% 保留覆盖率下的归一化预测误差。",
            },
            {
                "sheet": "coverage_plot_summary",
                "description": "上图直接绘制的数据：分类/回归端点间中位数与 10,000 次 bootstrap 95% 区间。",
            },
            {
                "sheet": "trusted_gain_endpoints",
                "description": "下图直接绘制的 22 个点：各端点最高相似度四分位的误差下降百分比。",
            },
            {
                "sheet": "trusted_gain_summary",
                "description": "下图菱形和粗线：中位数、四分位数及改善端点数。",
            },
        ]
    )
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        molecule_data.to_excel(writer, sheet_name="raw_molecule_data", index=False)
        coverage.to_excel(writer, sheet_name="coverage_by_endpoint", index=False)
        coverage_summary.to_excel(writer, sheet_name="coverage_plot_summary", index=False)
        gains.to_excel(writer, sheet_name="trusted_gain_endpoints", index=False)
        gain_summary.to_excel(writer, sheet_name="trusted_gain_summary", index=False)
    format_workbook(workbook_path)

    print(f"Wrote {workbook_path}")
    print(f"Coverage endpoint rows: {len(coverage)}")
    print(f"Coverage plotted rows: {len(coverage_summary)}")
    print(f"Gain endpoint rows: {len(gains)}")


if __name__ == "__main__":
    main()
