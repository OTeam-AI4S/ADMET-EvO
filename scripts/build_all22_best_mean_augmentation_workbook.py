"""Build the final 22-task best-mean augmentation audit workbook."""

from __future__ import annotations

import json
import math
import statistics
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
MAPPING = ROOT / "runs/v27-maplight-gnn-freeze/full-22-task.mapping.json"
SCORECARD = ROOT / "reports/goal_scorecard_v27.json"
REPORT_ROOT = ROOT / "reports/all22_optimal_recipe_adaptive_data_v2"
RAW_DEVELOPMENT = REPORT_ROOT / "development"
DEVELOPMENT = REPORT_ROOT / "development_best_mean"
CONFIRMATION = REPORT_ROOT / "confirmation_best_mean"
EARLIER_CONFIRMATION = REPORT_ROOT / "confirmation"
OUTPUT = REPORT_ROOT / "all22_best_mean_augmentation_audit.xlsx"


def _load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _artifact(directory: Path, task: str) -> dict[str, Any] | None:
    path = directory / f"{task}.json"
    return _load(path) if path.exists() else None


def _recipe_label(frozen: dict[str, Any]) -> str:
    if frozen.get("kind") == "recipe":
        return f"recipe({frozen['recipe']})"
    if frozen.get("kind") == "ensemble":
        return "ensemble(" + "+".join(row["model"] for row in frozen["members"]) + ")"
    return str(frozen["model"])


def _sample_variance_std(values: list[float]) -> tuple[float | None, float | None]:
    if not values:
        return None, None
    variance = statistics.variance(values) if len(values) > 1 else 0.0
    return variance, math.sqrt(variance)


def _matching_confirmation(task: str, ratio: float | None) -> dict[str, Any] | None:
    if ratio is None:
        return None
    for directory in (CONFIRMATION, EARLIER_CONFIRMATION):
        payload = _artifact(directory, task)
        selected = payload.get("selected_ratio") if payload else None
        if selected is not None and math.isclose(
            float(selected), ratio, rel_tol=0.0, abs_tol=1e-12
        ):
            return payload
    return None


def _style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_table(sheet, name: str, last_column: str) -> None:
    table = Table(displayName=name, ref=f"A1:{last_column}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def main() -> None:
    mapping = _load(MAPPING)
    scorecard_rows = _load(SCORECARD)["tasks"]
    scorecard = {row["task"]: row for row in scorecard_rows}
    main_rows: list[list[Any]] = []
    sweep_rows: list[list[Any]] = []

    for score_row in scorecard_rows:
        task = score_row["task"]
        development = _artifact(DEVELOPMENT, task)
        best = development.get("best_mean_selection") if development else None
        ratio = float(best["ratio"]) if best else None
        confirmation = _matching_confirmation(task, ratio)

        if best:
            dev_baseline = best.get("baseline_mean")
            dev_augmented = best.get("augmented_mean")
            dev_delta = best.get("oriented_delta_mean")
            dev_variance = best.get("oriented_delta_sample_variance")
            dev_sd = best.get("oriented_delta_sample_standard_deviation")
            dev_positive = best.get("positive_seed_fraction")
        else:
            dev_baseline = dev_augmented = dev_delta = None
            dev_variance = dev_sd = dev_positive = None

        if confirmation:
            summary = confirmation["agent_summary"]
            confirmation_deltas = [
                float(row["oriented_delta"])
                for row in confirmation.get("per_seed_model", [])
            ]
            conf_variance, conf_sd = _sample_variance_std(confirmation_deltas)
            final_status = (
                "best_ratio_confirmation_positive"
                if float(summary["oriented_delta_mean"]) > 0
                else "retain_baseline_confirmation_nonpositive"
            )
        else:
            summary = {}
            conf_variance = conf_sd = None
            if not development:
                final_status = "development_pending"
            elif best is None:
                final_status = development.get("status", "no_ratio_sweep")
            elif float(dev_delta) <= 0:
                final_status = "retain_baseline_best_development_nonpositive"
            else:
                final_status = "best_ratio_confirmation_pending"

        main_rows.append(
            [
                task,
                score_row["metric"],
                _recipe_label(mapping[task]),
                float(scorecard[task]["score"]),
                dev_baseline,
                ratio,
                dev_augmented,
                dev_delta,
                dev_variance,
                dev_sd,
                dev_positive,
                summary.get("baseline_mean"),
                summary.get("augmented_mean"),
                summary.get("oriented_delta_mean"),
                conf_variance,
                conf_sd,
                summary.get("positive_seed_fraction"),
                final_status,
            ]
        )

        if development:
            for ratio_text, ratio_summary in development.get("ratio_summaries", {}).items():
                sweep_ratio = float(ratio_text)
                deltas = [
                    float(row["oriented_delta"])
                    for row in development.get("per_seed_ratio_model", [])
                    if math.isclose(
                        float(row["ratio"]),
                        sweep_ratio,
                        rel_tol=0.0,
                        abs_tol=1e-12,
                    )
                ]
                variance, sd = _sample_variance_std(deltas)
                sweep_rows.append(
                    [
                        task,
                        score_row["metric"],
                        sweep_ratio,
                        ratio_summary.get("baseline_mean"),
                        ratio_summary.get("augmented_mean"),
                        ratio_summary.get("oriented_delta_mean"),
                        variance,
                        sd,
                        ratio_summary.get("positive_seed_fraction"),
                        ratio is not None
                        and math.isclose(
                            sweep_ratio, ratio, rel_tol=0.0, abs_tol=1e-12
                        ),
                    ]
                )

    workbook = Workbook()
    results = workbook.active
    results.title = "Best mean 22-task audit"
    results.append(
        [
            "Task",
            "Metric",
            "Frozen v2.7 Recipe",
            "v2.7 Official Score (reference only)",
            "Dev Baseline Mean",
            "Best Ratio",
            "Dev Augmented Mean",
            "Dev Oriented Δ Mean",
            "Dev Δ Sample Variance",
            "Dev Δ SD",
            "Dev Positive Seeds",
            "Confirmation Baseline",
            "Confirmation Augmented",
            "Confirmation Oriented Δ Mean",
            "Confirmation Δ Sample Variance",
            "Confirmation Δ SD",
            "Confirmation Positive Seeds",
            "Final Status",
        ]
    )
    for row in main_rows:
        results.append(row)
    results.freeze_panes = "A2"
    results.row_dimensions[1].height = 42
    _style_header(results)
    widths = [34, 11, 52, 22, 20, 14, 21, 21, 22, 18, 18, 22, 22, 24, 25, 18, 22, 45]
    for index, width in enumerate(widths, start=1):
        results.column_dimensions[results.cell(1, index).column_letter].width = width
    for row in results.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in range(4, 18):
            row[index - 1].number_format = "0.000000"
    _add_table(results, "BestMeanAugmentationAudit", "R")
    for column in ("H", "N"):
        results.conditional_formatting.add(
            f"{column}2:{column}{results.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="num",
                mid_value=0,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    sweep = workbook.create_sheet("All ratio sweep")
    sweep.append(
        [
            "Task",
            "Metric",
            "Ratio",
            "Baseline Mean",
            "Augmented Mean",
            "Oriented Δ Mean",
            "Δ Sample Variance",
            "Δ SD",
            "Positive Seeds",
            "Selected Best Ratio",
        ]
    )
    for row in sweep_rows:
        sweep.append(row)
    sweep.freeze_panes = "A2"
    _style_header(sweep)
    for column, width in zip(
        "ABCDEFGHIJ", [34, 11, 12, 20, 20, 20, 22, 18, 18, 20], strict=True
    ):
        sweep.column_dimensions[column].width = width
    for row in sweep.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in range(3, 10):
            row[index - 1].number_format = "0.000000"
    if sweep.max_row > 1:
        _add_table(sweep, "AllRatioSweep", "J")
        sweep.conditional_formatting.add(
            f"F2:F{sweep.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="num",
                mid_value=0,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    notes = workbook.create_sheet("Protocol notes")
    notes_rows = [
        ("Item", "Value"),
        ("Comparison", "Within each task and seed, baseline and augmented arms use the identical frozen v2.7 recipe; only admitted training membership changes."),
        ("Best-ratio selection", "Choose the ratio with maximum mean oriented delta over Development seeds 2-5. Exact ties choose the smaller ratio."),
        ("Variance", "Variance columns are sample variance across four seeds (ddof=1); SD is its square root."),
        ("Development", "TDC default scaffold seeds 2-5; full 0.5x, 1x, 2x, 4x and 8x sweep."),
        ("Confirmation", "Disjoint seeds 8-11 at the Development best-mean ratio for positive-mean candidates."),
        ("Official score caveat", "Official scores are reference-only test results and are not subtracted from scaffold-validation scores as a paired effect."),
        ("Runtime", "Slurm job 172620 on compute-0002; four A100 GPUs for Uni-Mol. MapLight uses isolated torch 2.2.1+cpu, DGL 2.1.0 and molfeat 0.11.0."),
        ("Frozen mapping", str(MAPPING.relative_to(ROOT))),
        ("Raw Development", str(RAW_DEVELOPMENT.relative_to(ROOT))),
        ("Best-mean Development", str(DEVELOPMENT.relative_to(ROOT))),
        ("Best-mean Confirmation", str(CONFIRMATION.relative_to(ROOT))),
    ]
    for row in notes_rows:
        notes.append(row)
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 120
    _style_header(notes)
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
