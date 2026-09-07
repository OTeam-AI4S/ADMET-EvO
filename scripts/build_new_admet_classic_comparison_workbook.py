"""Build the review-ready Excel workbook for the 17-task classic-model panel."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "2F8F9D"
GREEN = "70AD47"
GOLD = "FFC000"
RED = "C00000"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_GOLD = "FFF2CC"
LIGHT_RED = "FCE4D6"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="D9E1F2")


CLASSICS = {
    "ExtraTrees-Morgan": "baseline",
    "ExtraTrees-rich": "feature_rich",
    "LightGBM-fixed": "boosting_anchor",
    "LightGBM-tuned": "automl_expanded",
    "CatBoost-GPU": "gpu_catboost",
    "RandomForest-rich": "random_forest",
    "LinearSVM-Morgan": "linear_svm",
    "XGBoost-rich": "xgboost",
}


def _native(value: float, metric: str) -> float:
    return -float(value) if metric == "mae" else float(value)


def _style_table(ws, *, freeze: str = "A2", filter_table: bool = True) -> None:
    ws.freeze_panes = freeze
    if filter_table and ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=GRID)


def _set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _append_frame(ws, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))


def _comparison(root: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    expanded = root / "reports/new_admet_v5_expanded_actions_summary_v2"
    summary = pd.read_csv(expanded / "endpoint_confirmation.csv")
    folds = pd.read_csv(expanded / "fold_results.csv")
    classic_files = sorted(
        (root / "runs/new_admet/classic_panel_17_v1").glob("*.csv")
    )
    if len(classic_files) != 34:
        raise RuntimeError(f"Expected 34 classic-panel shards, found {len(classic_files)}")
    classic_panel = pd.concat(
        [pd.read_csv(path) for path in classic_files], ignore_index=True
    )
    hardness = pd.read_csv(root / "reports/new_admet_v5_hardness_v1.csv")
    confirm = folds.loc[folds["confirmation_seed"].isin([29, 43])].copy()
    pivot = (
        confirm.groupby(["endpoint_id", "action"], as_index=False)["oriented_score"]
        .mean()
        .pivot(index="endpoint_id", columns="action", values="oriented_score")
    )
    rows = []
    oriented_rows = []
    hardness_map = hardness.set_index("endpoint_id")["hardness"].to_dict()
    for record in summary.itertuples(index=False):
        metric = str(record.metric)
        agent_oriented = float(record.confirmation_score)
        classic_oriented = {
            label: float(pivot.loc[record.endpoint_id, action])
            for label, action in CLASSICS.items()
        }
        best_label, best_oriented = max(
            classic_oriented.items(), key=lambda item: (item[1], item[0])
        )
        difference = agent_oriented - best_oriented
        status = "胜" if difference > 1e-12 else "负" if difference < -1e-12 else "平"
        all_oriented = {"Agent": agent_oriented, **classic_oriented}
        agent_rank = (
            pd.Series(all_oriented).rank(ascending=False, method="average").loc["Agent"]
        )
        rows.append({
            "Endpoint": record.endpoint_id,
            "Source": record.source_id,
            "Task type": record.task_type,
            "Metric": "MCC ↑" if metric == "mcc" else "MAE ↓",
            "Agent selected action": record.selected_action,
            "Agent score": _native(agent_oriented, metric),
            **{label: _native(value, metric) for label, value in classic_oriented.items()},
            "Best classic model": best_label,
            "Best classic score": _native(best_oriented, metric),
            "Agent vs best classic": status,
            "Oriented delta vs best": difference,
            f"Agent rank / {len(CLASSICS) + 1}": float(agent_rank),
            "Seed SD": float(record.confirmation_seed_sd),
            "Hardness": float(hardness_map.get(record.endpoint_id, float("nan"))),
        })
        oriented_rows.append({"Endpoint": record.endpoint_id, **all_oriented})
    comparison = pd.DataFrame(rows).sort_values(
        ["Task type", "Endpoint"], kind="stable"
    ).reset_index(drop=True)
    oriented = pd.DataFrame(oriented_rows).set_index("Endpoint")
    wtl = []
    for label in CLASSICS:
        delta = oriented["Agent"] - oriented[label]
        wtl.append({
            "Comparator": label,
            "Agent wins": int(delta.gt(1e-12).sum()),
            "Ties": int(delta.abs().le(1e-12).sum()),
            "Agent losses": int(delta.lt(-1e-12).sum()),
            "Median oriented gain": float(delta.median()),
            "Mean oriented gain": float(delta.mean()),
            "Interpretation": (
                "明显超过" if int(delta.gt(1e-12).sum()) >= 12
                else "基本相当" if abs(int(delta.gt(1e-12).sum()) - int(delta.lt(-1e-12).sum())) <= 2
                else "略优" if int(delta.gt(1e-12).sum()) > int(delta.lt(-1e-12).sum())
                else "落后"
            ),
        })
    return comparison, pd.DataFrame(wtl)


def build(root: Path, output: Path) -> None:
    comparison, wtl = _comparison(root)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "17 New ADMET Tasks — Classic Model Comparison"
    workbook.properties.subject = "Matched chemistry-group holdout benchmark"
    workbook.properties.creator = "ADMET Agent"

    overview = workbook.create_sheet("Overview")
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:H1")
    overview["A1"] = "17个新ADMET任务：Agent与经典模型公平对比"
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].font = Font(color=WHITE, bold=True, size=16)
    overview["A1"].alignment = Alignment(horizontal="center", vertical="center")
    overview.row_dimensions[1].height = 32
    overview.append([])
    overview.append(["结论", "结果"])
    for _, record in wtl.iterrows():
        overview.append([
            f"Agent vs {record['Comparator']}",
            f"{record['Agent wins']}胜 / {record['Ties']}平 / "
            f"{record['Agent losses']}负；{record['Interpretation']}",
        ])
    overview.append(["总体定位", "以相同分子、相同化学分组切分和相同确认种子做匹配比较"])
    overview.append([])
    evidence_header_row = overview.max_row + 1
    overview.append(["证据边界", "说明"])
    overview.append(["切分", "完全相同的5个identity/chemistry-group locked folds"])
    overview.append(["确认种子", "29、43；种子17仅用于Agent action选择"])
    overview.append(["指标", "分类=MCC（越高越好）；回归=MAE（越低越好）"])
    overview.append(["论文数字", "仅作背景参考；不同split/metric不可直接判定胜负"])
    overview.append(["50→20方案", "允许探索，但必须公开50-task全分母，并在未触碰confirmation上确认Top-20"])
    for cell in overview[3]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
    for cell in overview[evidence_header_row]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
    for row in overview.iter_rows(min_row=4, max_row=3 + len(wtl) + 1):
        row[0].font = Font(bold=True, color=NAVY)
        row[1].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    for row in overview.iter_rows(min_row=evidence_header_row + 1, max_row=evidence_header_row + 5):
        row[0].font = Font(bold=True, color=NAVY)
        row[1].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    _set_widths(overview, {"A": 31, "B": 84, "C": 3, "D": 3, "E": 3, "F": 3, "G": 3, "H": 3})
    overview.freeze_panes = "A3"

    task_sheet = workbook.create_sheet("17 Task Comparison")
    _append_frame(task_sheet, comparison)
    _style_table(task_sheet)
    for column in range(1, task_sheet.max_column + 1):
        task_sheet.column_dimensions[get_column_letter(column)].width = 17
    task_sheet.column_dimensions["A"].width = 42
    task_sheet.column_dimensions["B"].width = 28
    task_sheet.column_dimensions["E"].width = 23
    headers = {cell.value: cell.column for cell in task_sheet[1]}
    for row in range(2, task_sheet.max_row + 1):
        for column in range(6, task_sheet.max_column + 1):
            task_sheet.cell(row, column).number_format = "0.000"
        status = task_sheet.cell(row, headers["Agent vs best classic"])
        status.fill = PatternFill(
            "solid", fgColor=LIGHT_GREEN if status.value == "胜" else LIGHT_RED if status.value == "负" else LIGHT_GOLD
        )
        status.font = Font(bold=True, color=GREEN if status.value == "胜" else RED if status.value == "负" else NAVY)
    task_sheet.conditional_formatting.add(
        f"{get_column_letter(headers['Oriented delta vs best'])}2:{get_column_letter(headers['Oriented delta vs best'])}{task_sheet.max_row}",
        ColorScaleRule(start_type="min", start_color=RED, mid_type="percentile", mid_value=50,
                       mid_color=LIGHT_GOLD, end_type="max", end_color=GREEN),
    )
    task_sheet.conditional_formatting.add(
        f"{get_column_letter(headers[f'Agent rank / {len(CLASSICS) + 1}'])}2:{get_column_letter(headers[f'Agent rank / {len(CLASSICS) + 1}'])}{task_sheet.max_row}",
        ColorScaleRule(start_type="min", start_color=GREEN, mid_type="percentile", mid_value=50,
                       mid_color=LIGHT_GOLD, end_type="max", end_color=RED),
    )

    wtl_sheet = workbook.create_sheet("Win-Tie-Loss")
    _append_frame(wtl_sheet, wtl)
    _style_table(wtl_sheet)
    _set_widths(wtl_sheet, {"A": 24, "B": 14, "C": 11, "D": 14, "E": 23, "F": 22, "G": 18})
    for row in range(2, wtl_sheet.max_row + 1):
        wtl_sheet.cell(row, 5).number_format = "0.000"
        wtl_sheet.cell(row, 6).number_format = "0.000"
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Agent vs classic models (17 endpoints)"
    chart.y_axis.title = "Comparator"
    chart.x_axis.title = "Endpoint count"
    data = Reference(wtl_sheet, min_col=2, max_col=4, min_row=1, max_row=wtl_sheet.max_row)
    categories = Reference(wtl_sheet, min_col=1, min_row=2, max_row=wtl_sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 15
    wtl_sheet.add_chart(chart, "I2")

    chart_data = workbook.create_sheet("Chart Data")
    score_columns = ["Agent score", *CLASSICS]
    classification = comparison.loc[
        comparison["Task type"].eq("classification"),
        ["Endpoint", *score_columns, "Best classic score", "Best classic model"],
    ].copy()
    classification["Endpoint"] = classification["Endpoint"].str.replace(
        r"^(tox21::|b3db::|clintox::|ncats_pampa::)", "", regex=True
    )
    regression = comparison.loc[
        comparison["Task type"].eq("regression"),
        ["Endpoint", *score_columns, "Best classic score", "Best classic model"],
    ].copy()
    regression["Endpoint"] = regression["Endpoint"].str.replace(
        r"^(b3db::|freesolv::)", "", regex=True
    )
    model_rank_rows = []
    for _, record in comparison.iterrows():
        metric = record["Metric"]
        values = {"Agent": record["Agent score"], **{
            label: record[label] for label in CLASSICS
        }}
        model_rank_rows.append(
            pd.Series(values).rank(
                ascending=metric == "MAE ↓", method="average"
            )
        )
    mean_ranks = pd.DataFrame(model_rank_rows).mean().sort_values().reset_index()
    mean_ranks.columns = ["Model", "Mean rank (lower is better)"]

    chart_data["A1"] = "Classification MCC — all models"
    all_score_count = len(score_columns) + 1
    for column, name in enumerate(classification.columns[:all_score_count], start=1):
        chart_data.cell(2, column, name)
    for row_index, row in enumerate(classification.iloc[:, :all_score_count].itertuples(index=False, name=None), start=3):
        for column, value in enumerate(row, start=1):
            chart_data.cell(row_index, column, value)
    start_cls_best = all_score_count + 3
    chart_data.cell(1, start_cls_best, "Classification MCC — Agent vs best classic")
    for offset, name in enumerate(["Endpoint", "Agent", "Best classic", "Best classic model"]):
        chart_data.cell(2, start_cls_best + offset, name)
    for row_index, row in enumerate(classification.itertuples(index=False), start=3):
        chart_data.cell(row_index, start_cls_best, row[0])
        chart_data.cell(row_index, start_cls_best + 1, row[1])
        chart_data.cell(row_index, start_cls_best + 2, row[-2])
        chart_data.cell(row_index, start_cls_best + 3, row[-1])

    start_regression = start_cls_best + 6
    chart_data.cell(1, start_regression, "Regression MAE — all models")
    for offset, name in enumerate(regression.columns[:all_score_count]):
        chart_data.cell(2, start_regression + offset, name)
    for row_index, row in enumerate(regression.iloc[:, :all_score_count].itertuples(index=False, name=None), start=3):
        for offset, value in enumerate(row):
            chart_data.cell(row_index, start_regression + offset, value)

    start_rank = start_regression + all_score_count + 2
    chart_data.cell(1, start_rank, "Mean model rank")
    for offset, name in enumerate(mean_ranks.columns):
        chart_data.cell(2, start_rank + offset, name)
    for row_index, row in enumerate(mean_ranks.itertuples(index=False, name=None), start=3):
        chart_data.cell(row_index, start_rank, row[0])
        chart_data.cell(row_index, start_rank + 1, float(row[1]))
    _style_table(chart_data, freeze="A3", filter_table=False)
    for merged_start in (1, start_cls_best, start_regression, start_rank):
        cell = chart_data.cell(1, merged_start)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
    for column in range(2, chart_data.max_column + 1):
        chart_data.column_dimensions[get_column_letter(column)].width = 18
    chart_data.column_dimensions["A"].width = 33
    chart_data.column_dimensions["J"].width = 33
    chart_data.column_dimensions["O"].width = 43
    chart_data.column_dimensions["W"].width = 24
    for row in chart_data.iter_rows(min_row=3):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    charts_sheet = workbook.create_sheet("Numeric Charts")
    charts_sheet.sheet_view.showGridLines = False
    charts_sheet["A1"] = "17-task numerical model comparison"
    charts_sheet["A1"].font = Font(size=16, bold=True, color=NAVY)

    cls_all = BarChart()
    cls_all.type = "bar"
    cls_all.style = 10
    cls_all.grouping = "clustered"
    cls_all.title = "Classification MCC — all matched models"
    cls_all.x_axis.title = "MCC (higher is better)"
    cls_all.y_axis.title = "Endpoint"
    cls_all.height = 14
    cls_all.width = 28
    cls_all.add_data(
        Reference(chart_data, min_col=2, max_col=all_score_count, min_row=2, max_row=2 + len(classification)),
        titles_from_data=True,
    )
    cls_all.set_categories(
        Reference(chart_data, min_col=1, min_row=3, max_row=2 + len(classification))
    )
    charts_sheet.add_chart(cls_all, "A3")

    cls_best = BarChart()
    cls_best.type = "bar"
    cls_best.style = 11
    cls_best.grouping = "clustered"
    cls_best.title = "Classification MCC — Agent vs best classic per endpoint"
    cls_best.x_axis.title = "MCC (higher is better)"
    cls_best.y_axis.title = "Endpoint"
    cls_best.height = 14
    cls_best.width = 22
    cls_best.add_data(
        Reference(
            chart_data,
            min_col=start_cls_best + 1,
            max_col=start_cls_best + 2,
            min_row=2,
            max_row=2 + len(classification),
        ),
        titles_from_data=True,
    )
    cls_best.set_categories(
        Reference(
            chart_data,
            min_col=start_cls_best,
            min_row=3,
            max_row=2 + len(classification),
        )
    )
    cls_best.dLbls = DataLabelList()
    cls_best.dLbls.showVal = True
    charts_sheet.add_chart(cls_best, "A31")

    reg_all = BarChart()
    reg_all.type = "col"
    reg_all.style = 12
    reg_all.grouping = "clustered"
    reg_all.title = "Regression MAE — all matched models"
    reg_all.y_axis.title = "MAE (lower is better)"
    reg_all.x_axis.title = "Endpoint"
    reg_all.height = 9
    reg_all.width = 22
    reg_all.add_data(
        Reference(
            chart_data,
            min_col=start_regression + 1,
            max_col=start_regression + all_score_count - 1,
            min_row=2,
            max_row=2 + len(regression),
        ),
        titles_from_data=True,
    )
    reg_all.set_categories(
        Reference(
            chart_data,
            min_col=start_regression,
            min_row=3,
            max_row=2 + len(regression),
        )
    )
    reg_all.dLbls = DataLabelList()
    reg_all.dLbls.showVal = True
    charts_sheet.add_chart(reg_all, "X3")

    rank_chart = BarChart()
    rank_chart.type = "bar"
    rank_chart.style = 13
    rank_chart.title = "Mean rank across 17 endpoints"
    rank_chart.x_axis.title = "Mean rank (lower is better)"
    rank_chart.y_axis.title = "Model"
    rank_chart.height = 8
    rank_chart.width = 16
    rank_chart.add_data(
        Reference(
            chart_data,
            min_col=start_rank + 1,
            min_row=2,
            max_row=2 + len(mean_ranks),
        ),
        titles_from_data=True,
    )
    rank_chart.set_categories(
        Reference(
            chart_data,
            min_col=start_rank,
            min_row=3,
            max_row=2 + len(mean_ranks),
        )
    )
    rank_chart.dLbls = DataLabelList()
    rank_chart.dLbls.showVal = True
    charts_sheet.add_chart(rank_chart, "X22")
    _set_widths(charts_sheet, {get_column_letter(index): 12 for index in range(1, start_rank + 8)})

    literature = pd.DataFrame([
        {
            "Dataset": "Tox21 (12 endpoints)",
            "Published model/result": "Kernel SVM ROC-AUC 0.822; graph model 0.829",
            "Our result": "MCC median ≈0.300 across 12 endpoints",
            "Directly comparable?": "No",
            "Reason": "Published random split + ROC-AUC; ours chemistry-group split + MCC",
            "Source": "https://pmc.ncbi.nlm.nih.gov/articles/PMC5868307/",
        },
        {
            "Dataset": "ClinTox",
            "Published model/result": "Bypass/Weave ROC-AUC ≈0.827–0.832",
            "Our result": "MCC 0.052",
            "Directly comparable?": "No",
            "Reason": "Published random split + ROC-AUC; ours chemistry-group split + MCC",
            "Source": "https://pmc.ncbi.nlm.nih.gov/articles/PMC5868307/",
        },
        {
            "Dataset": "FreeSolv",
            "Published model/result": "XGBoost RMSE 1.74; MPNN RMSE 1.15",
            "Our result": "MAE 2.273",
            "Directly comparable?": "No",
            "Reason": "RMSE vs MAE and random vs chemistry-group split",
            "Source": "https://pmc.ncbi.nlm.nih.gov/articles/PMC5868307/",
        },
        {
            "Dataset": "B3DB BBB classification",
            "Published model/result": "LightBBB MCC 0.43; Deep-B 0.49; cliques 0.74",
            "Our result": "MCC 0.573",
            "Directly comparable?": "Reference only",
            "Reason": "Different curation, retained molecules and validation splits",
            "Source": "https://pubs.acs.org/doi/10.1021/acs.jcim.4c01217",
        },
        {
            "Dataset": "B3DB logBB",
            "Published model/result": "Literature usually reports RMSE/R²",
            "Our result": "MAE 0.445",
            "Directly comparable?": "No",
            "Reason": "Different metrics and splits",
            "Source": "https://github.com/theochem/B3DB",
        },
        {
            "Dataset": "NCATS PAMPA",
            "Published model/result": "No single standardized public leaderboard found",
            "Our result": "MCC 0.248",
            "Directly comparable?": "No",
            "Reason": "Use matched re-evaluation instead",
            "Source": "https://opendata.ncats.nih.gov/adme/",
        },
    ])
    literature_sheet = workbook.create_sheet("Literature Reference")
    _append_frame(literature_sheet, literature)
    _style_table(literature_sheet)
    _set_widths(literature_sheet, {"A": 28, "B": 48, "C": 35, "D": 20, "E": 58, "F": 62})
    for row in range(2, literature_sheet.max_row + 1):
        cell = literature_sheet.cell(row, 6)
        cell.hyperlink = cell.value
        cell.style = "Hyperlink"
        if literature_sheet.cell(row, 4).value == "No":
            literature_sheet.cell(row, 4).fill = PatternFill("solid", fgColor=LIGHT_RED)
        else:
            literature_sheet.cell(row, 4).fill = PatternFill("solid", fgColor=LIGHT_GOLD)

    notes = pd.DataFrame([
        ["Agent score", "Per-endpoint action, including RF/SVM/XGBoost warm starts, selected only on seed 17, then evaluated on confirmation seeds 29 and 43."],
        ["Classic panel", "Every model uses exactly the same curated molecules and identity/chemistry-group locked folds."],
        ["ExtraTrees-Morgan", "200-tree ExtraTrees on Morgan radius-2 fingerprints."],
        ["ExtraTrees-rich", "Same ExtraTrees family on Morgan + Avalon + ErG + RDKit2D."],
        ["LightGBM-fixed", "Fixed 250-tree LightGBM anchor; no outer-fold tuning."],
        ["LightGBM-tuned", "Four fixed candidates selected by group-disjoint inner CV only."],
        ["CatBoost-GPU", "600-iteration CatBoost on rich 2D features, threshold fixed at 0.5 for classification."],
        ["RandomForest-rich", "400-tree class-balanced RandomForest on Morgan + Avalon + ErG + RDKit2D."],
        ["LinearSVM-Morgan", "Fixed C=1 linear SVM on scaled Morgan fingerprints; class-balanced for classification."],
        ["XGBoost-rich", "Fixed 500-tree histogram XGBoost on rich 2D features; class imbalance weighted for classification."],
        ["Oriented delta", "All metrics are oriented so positive means Agent is better; regression MAE is negated only for comparison."],
        ["Literature", "Published scores are contextual references and must not be used in the matched W/T/L totals."],
        ["Top-20 policy", "If screening 50 tasks, retain the full 50-task denominator and confirm the selected 20 on untouched data."],
    ], columns=["Item", "Definition / constraint"])
    notes_sheet = workbook.create_sheet("Methods & Notes")
    _append_frame(notes_sheet, notes)
    _style_table(notes_sheet, filter_table=False)
    _set_widths(notes_sheet, {"A": 24, "B": 110})

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path("."))
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("reports/new_admet_17task_classic_model_comparison_v2_2026-08-27.xlsx"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.root.resolve(), args.output.resolve())
