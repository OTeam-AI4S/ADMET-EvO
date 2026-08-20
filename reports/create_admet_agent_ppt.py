#!/usr/bin/env python3
"""Create a concise two-slide ADMET Agent V2.1 / Uni-Mol2 briefing."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/wyilin/admet")
LEADERBOARD = ROOT / "reports/leaderboard_snapshot_2026-08-13.json"
COMPARISON = ROOT / "runs/leaderboard-v21/comparison.json"
OUTPUT = ROOT / "reports/ADMET_Agent_V2.1_UniMol2_2slides_EN.pptx"

FONT = "Arial"
NAVY = RGBColor(15, 28, 50)
NAVY_2 = RGBColor(27, 47, 78)
INK = RGBColor(26, 37, 55)
MUTED = RGBColor(94, 108, 129)
WHITE = RGBColor(255, 255, 255)
OFFWHITE = RGBColor(247, 249, 252)
LIGHT_BLUE = RGBColor(231, 240, 249)
GRID = RGBColor(211, 220, 231)
TEAL = RGBColor(28, 151, 146)
GREEN = RGBColor(36, 151, 95)
ORANGE = RGBColor(238, 143, 55)
RED = RGBColor(207, 72, 72)
PURPLE = RGBColor(110, 91, 181)


def set_run(run, text, size, color=INK, bold=False, font=FONT):
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text_box(slide, x, y, w, h, text, size=12, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.03):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold)
    return shape


def rounded_box(slide, x, y, w, h, fill, title, body="", accent=None,
                title_size=13, body_size=9.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = accent or fill
    shape.line.width = Pt(1.1)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.11)
    tf.margin_top = tf.margin_bottom = Inches(0.07)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), title, title_size, WHITE if fill in (NAVY, NAVY_2, TEAL, PURPLE) else INK, True)
    if body:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(3)
        set_run(p2.add_run(), body, body_size, WHITE if fill in (NAVY, NAVY_2, TEAL, PURPLE) else MUTED)
    return shape


def add_title(slide, title, subtitle, slide_no):
    text_box(slide, 0.48, 0.24, 10.8, 0.42, title, 23, NAVY, True)
    text_box(slide, 0.5, 0.67, 11.6, 0.25, subtitle, 9.2, MUTED)
    text_box(slide, 12.35, 0.25, 0.48, 0.32, f"0{slide_no}", 12, TEAL, True, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.99), Inches(12.33), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = GRID; line.line.fill.background()


def add_arrow(slide, x, y, w=0.24, h=0.30):
    arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = TEAL
    arrow.line.fill.background()


def metric_text(metric, direction):
    names = {"auroc": "AUROC", "auprc": "AUPRC", "mae": "MAE", "spearman": "Spearman"}
    return names[metric] + (" ↓" if direction == "minimize" else " ↑")


def fmt(v):
    return f"{v:.3f}" if abs(v) < 10 else f"{v:.3f}"


def rank_text(row):
    best = row["estimated_rank"]
    worst = row.get("tie_worst_rank", best)
    n = row["existing_entries"] + 1
    prefix = str(best) if best == worst else f"{best}–{worst}"
    return f"{prefix}/{n}"


def solution_text(finalist):
    names = {
        "extra_trees": "ET", "xgboost": "XGB", "random_forest": "RF",
        "svm_rbf": "RBF-SVM", "lightgbm": "LGB"
    }
    if finalist.startswith("ensemble("):
        inner = finalist[len("ensemble("):-1]
        models = "+".join(names.get(x, x) for x in inner.split("+"))
    else:
        models = names.get(finalist, finalist)
    return f"Multi-view 2D (Morgan/Avalon/ErG/RDKit2D) → {models}"


def style_cell(cell, text, size, color=INK, bold=False, fill=None,
               align=PP_ALIGN.CENTER, left=0.04):
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(left)
    cell.margin_top = cell.margin_bottom = Inches(0.012)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold)


def slide_one(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = OFFWHITE
    add_title(
        slide,
        "ADMET Agent: Endpoint-Adaptive 2D–3D Evidence Routing",
        "The agent selects the most reliable expert for each endpoint based on data characteristics and uncertainty",
        1,
    )

    # Pipeline
    y = 1.24
    boxes = [
        (0.50, 1.50, "Task Input", "SMILES · endpoint\nmetric · budget", NAVY),
        (2.25, 1.65, "Data Governance", "scaffold split\nisolation audit", NAVY_2),
        (4.15, 1.68, "Endpoint Profile", "size · imbalance · 3D relevance\ncross-seed stability", PURPLE),
        (6.11, 2.18, "Expert Pool", "2D ML · Uni-Mol2 3D\ntarget transforms · ensembles", TEAL),
        (8.57, 1.74, "Evidence Router", "LOSO + LCB\nsingle · fusion · fallback", ORANGE),
        (10.55, 2.28, "Freeze & Evaluate", "lock recipe + hash\n5 seeds → fixed test", NAVY),
    ]
    for i, (x, w, title, body, color) in enumerate(boxes):
        rounded_box(slide, x, y, w, 0.94, color, title, body, title_size=12.2, body_size=8.6)
        if i < len(boxes) - 1:
            add_arrow(slide, x + w + 0.07, y + 0.32, 0.24, 0.30)

    # Key output badges
    text_box(slide, 0.52, 2.36, 3.4, 0.30, "FOUR INTERPRETABLE ROUTING DECISIONS", 10.5, NAVY, True)
    cards = [
        (0.52, "2D ONLY", "Bioavailability\nCYP2C9 Substrate", LIGHT_BLUE, NAVY),
        (3.10, "3D ONLY", "Lipophilicity\nUni-Mol2 clearly wins", RGBColor(226, 245, 241), GREEN),
        (5.68, "2D + 3D", "DILI · hERG\ncomplementary evidence", RGBColor(238, 235, 249), PURPLE),
        (8.26, "RANK FUSION", "Half-Life\n3D + rank targets", RGBColor(253, 239, 224), ORANGE),
    ]
    for x, title, body, fill, accent in cards:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.72), Inches(2.36), Inches(0.82))
        box.fill.solid(); box.fill.fore_color.rgb = fill
        box.line.color.rgb = accent; box.line.width = Pt(1.1)
        text_box(slide, x + 0.11, 2.79, 2.14, 0.22, title, 9.3, accent, True, PP_ALIGN.CENTER)
        text_box(slide, x + 0.11, 3.03, 2.14, 0.40, body, 8.3, INK, False, PP_ALIGN.CENTER)
    rounded_box(slide, 10.84, 2.72, 1.98, 0.82, NAVY, "V2.1 BASELINE", "Top-1  3/22\nTop-5  13/22", title_size=9.5, body_size=9)

    # Uni-Mol pilot table
    text_box(slide, 0.52, 3.78, 7.8, 0.32, "Uni-Mol2 84M: Strict Six-Task Scaffold Validation (Test Unused)", 11.2, NAVY, True)
    data = [
        ("Bioavailability", "AUROC ↑", "0.7146", "0.7008", "2D / cautious fusion"),
        ("CYP2C9 Sub.", "AUPRC ↑", "0.4801", "0.3563", "Fallback to 2D"),
        ("DILI", "AUROC ↑", "0.8654", "0.8635", "2D+3D fusion"),
        ("Half-Life", "Spearman ↑", "0.4374", "0.4511", "Heterogeneous rank fusion"),
        ("Lipophilicity", "MAE ↓", "0.5192", "0.4066", "Uni-Mol2"),
        ("hERG", "AUROC ↑", "0.8867", "0.8916", "2D+3D fusion"),
    ]
    tbl_shape = slide.shapes.add_table(7, 5, Inches(0.52), Inches(4.14), Inches(8.05), Inches(2.56))
    tbl = tbl_shape.table
    widths = [2.0, 1.20, 1.10, 1.25, 2.50]
    for col, width in zip(tbl.columns, widths): col.width = Inches(width)
    headers = ["Task", "Metric", "Best 2D", "Uni-Mol2", "Router Decision"]
    for j, h in enumerate(headers): style_cell(tbl.cell(0, j), h, 8.5, WHITE, True, NAVY)
    for i, row in enumerate(data, 1):
        fill = WHITE if i % 2 else LIGHT_BLUE
        for j, value in enumerate(row):
            color = GREEN if (j == 3 and row[0] in {"Lipophilicity", "hERG", "Half-Life"}) else INK
            bold = j == 4 or color == GREEN
            style_cell(tbl.cell(i, j), value, 8.2, color, bold, fill, PP_ALIGN.LEFT if j in (0, 4) else PP_ALIGN.CENTER)

    # Findings callout
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.82), Inches(3.92), Inches(4.0), Inches(2.78))
    box.fill.solid(); box.fill.fore_color.rgb = NAVY
    box.line.fill.background()
    text_box(slide, 9.05, 4.12, 3.55, 0.34, "THREE KEY FINDINGS", 13, WHITE, True)
    findings = [
        ("01", "Uni-Mol2 is an endpoint expert, not a universal replacement"),
        ("02", "84M provides a better score–compute trade-off than 164M"),
        ("03", "Mandatory fusion can hurt; the router must support fallback"),
    ]
    for k, (num, txt) in enumerate(findings):
        yy = 4.60 + k * 0.60
        text_box(slide, 9.06, yy, 0.42, 0.36, num, 10, TEAL, True, PP_ALIGN.CENTER)
        text_box(slide, 9.55, yy - 0.02, 2.98, 0.42, txt, 9.6, WHITE, k == 0)
    text_box(slide, 9.05, 6.40, 3.52, 0.18, "Next: freeze V2.3 recipes, then run the 22-task test once", 7.7, RGBColor(181, 196, 215), False)
    return slide


def slide_two(prs, leaderboard, comparison):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = OFFWHITE
    captured = leaderboard["captured_at"].replace("T", " ")[:16] + " UTC"
    add_title(
        slide,
        "ADMET Agent V2.1: 22-Task Fixed-Test Results vs. TDC Leaderboard",
        f"TDC public leaderboard snapshot: {captured} | Estimated rank after inserting our 5-seed mean into the public table",
        2,
    )

    badges = [
        (0.52, "TOP-1", "3 / 22", TEAL),
        (2.15, "TOP-3", "7 / 22", NAVY_2),
        (3.78, "TOP-5", "13 / 22", PURPLE),
        (5.41, "TOP-10", "19 / 22", ORANGE),
        (7.04, "MEDIAN RANK", "4.5", NAVY),
    ]
    for x, label, value, color in badges:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.16), Inches(1.45), Inches(0.54))
        box.fill.solid(); box.fill.fore_color.rgb = color; box.line.fill.background()
        text_box(slide, x + 0.08, 1.20, 0.76, 0.19, label, 6.8, WHITE, True)
        text_box(slide, x + 0.79, 1.20, 0.56, 0.21, value, 11.5, WHITE, True, PP_ALIGN.RIGHT)
    text_box(slide, 8.75, 1.19, 4.05, 0.36, "Projected Top-1: AMES · PPBR · Solubility", 9.3, GREEN, True, PP_ALIGN.RIGHT)

    rows = comparison["tasks"]
    page_map = leaderboard["pages"]
    table_shape = slide.shapes.add_table(23, 6, Inches(0.50), Inches(1.84), Inches(12.33), Inches(5.18))
    table = table_shape.table
    widths = [2.16, 0.88, 1.12, 0.72, 1.65, 5.80]
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    headers = ["Task", "Metric", "Our Score", "Rank", "TDC Top-3", "V2.1 Agent Solution (Validation-Selected)"]
    for j, h in enumerate(headers): style_cell(table.cell(0, j), h, 8.0, WHITE, True, NAVY, PP_ALIGN.LEFT if j in (0, 5) else PP_ALIGN.CENTER)

    short_names = {
        "CYP2C9_Substrate_CarbonMangels": "CYP2C9_Substrate",
        "CYP2D6_Substrate_CarbonMangels": "CYP2D6_Substrate",
        "CYP3A4_Substrate_CarbonMangels": "CYP3A4_Substrate",
        "Clearance_Hepatocyte_AZ": "Clearance_Hepatocyte",
        "Clearance_Microsome_AZ": "Clearance_Microsome",
        "Lipophilicity_AstraZeneca": "Lipophilicity",
        "Solubility_AqSolDB": "Solubility",
    }
    order = [
        "AMES", "BBB_Martins", "Bioavailability_Ma", "Caco2_Wang",
        "Clearance_Hepatocyte_AZ", "Clearance_Microsome_AZ",
        "CYP2C9_Substrate_CarbonMangels", "CYP2C9_Veith",
        "CYP2D6_Substrate_CarbonMangels", "CYP2D6_Veith",
        "CYP3A4_Substrate_CarbonMangels", "CYP3A4_Veith", "DILI",
        "Half_Life_Obach", "hERG", "HIA_Hou", "LD50_Zhu",
        "Lipophilicity_AstraZeneca", "Pgp_Broccatelli", "PPBR_AZ",
        "Solubility_AqSolDB", "VDss_Lombardo",
    ]
    by_task = {r["task"]: r for r in rows}
    for i, task in enumerate(order, 1):
        r = by_task[task]
        page = page_map[task]
        rank = rank_text(r)
        r_best = r["estimated_rank"]
        if r_best <= 3:
            fill = RGBColor(231, 246, 239)
            rank_color = GREEN
        elif r_best <= 5:
            fill = RGBColor(238, 242, 249)
            rank_color = NAVY_2
        elif r_best <= 10:
            fill = WHITE if i % 2 else LIGHT_BLUE
            rank_color = ORANGE
        else:
            fill = RGBColor(252, 239, 239)
            rank_color = RED
        top3 = " / ".join(fmt(x["mean"]) for x in page["entries"][:3])
        values = [
            short_names.get(task, task),
            metric_text(page["metric"], page["direction"]),
            f"{fmt(r['mean'])} ± {fmt(r['std'])}",
            rank,
            top3,
            solution_text(r["finalist"]),
        ]
        for j, value in enumerate(values):
            color = rank_color if j == 3 else INK
            bold = j == 3 or (j == 2 and r_best <= 3)
            align = PP_ALIGN.LEFT if j in (0, 5) else PP_ALIGN.CENTER
            style_cell(table.cell(i, j), value, 7.15 if j != 5 else 6.9, color, bold, fill, align)

    text_box(slide, 0.52, 7.10, 8.8, 0.20,
             "Note: Lower is better for MAE; higher is better otherwise. Rank ranges indicate ties; public scores are shown to three decimals.",
             7.2, MUTED)
    text_box(slide, 8.92, 7.10, 3.90, 0.20,
             "Not an official submission; new Uni-Mol2/router results are not included in these ranks",
             7.2, RED, True, PP_ALIGN.RIGHT)
    return slide


def main():
    if not LEADERBOARD.exists():
        raise FileNotFoundError(f"Missing leaderboard snapshot: {LEADERBOARD}")
    leaderboard = json.loads(LEADERBOARD.read_text())
    comparison = json.loads(COMPARISON.read_text())

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "ADMET Agent V2.1 and Uni-Mol2"
    prs.core_properties.subject = "Two-slide architecture and TDC leaderboard briefing"
    prs.core_properties.author = "ADMET Agent Team"
    slide_one(prs)
    slide_two(prs, leaderboard, comparison)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
